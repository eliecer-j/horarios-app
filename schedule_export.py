"""Exporta horarios en formato de grilla a Excel."""
from io import BytesIO

from datetime import timedelta

DAYS = ("Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom")


def build_schedule_workbook(start_date, end_date, people, assignments, branches, vacations):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:
        raise RuntimeError("Para descargar el Excel instala openpyxl: pip install openpyxl") from exc

    if end_date < start_date:
        raise ValueError("La fecha final debe ser igual o posterior a la inicial.")

    days = [start_date + timedelta(days=offset) for offset in range((end_date - start_date).days + 1)]
    if len(days) + 2 > 16384:
        raise ValueError("El rango supera el máximo de columnas que admite Excel.")

    assignments_by_day = {}
    for assignment in assignments:
        assigned_date, pid, branch_id, start_time, end_time = assignment[:5]
        assignment_type = assignment[8] if len(assignment) > 8 else "work"
        assignments_by_day[(assigned_date, pid)] = (
            branch_id, start_time, end_time, assignment_type
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Turnos"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B3"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:2"

    last_column = len(days) + 2
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title = sheet.cell(
        row=1, column=1,
        value=f"Horario de turnos · {start_date:%d/%m/%Y} al {end_date:%d/%m/%Y}",
    )
    title.fill = PatternFill("solid", fgColor="123B37")
    title.font = Font(name="Roboto", size=14, bold=True, color="FFFFFF")
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(name="Roboto", bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.cell(row=2, column=1, value="Persona")
    for column, assigned_date in enumerate(days, start=2):
        sheet.cell(
            row=2, column=column,
            value=f"{DAYS[assigned_date.weekday()]}\n{assigned_date:%d/%m/%Y}",
        )
    sheet.cell(row=2, column=last_column, value="Descansos")
    for cell in sheet[2]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    sheet.row_dimensions[2].height = 34

    grid_border = Border(
        right=Side(style="hair", color="CFD8D6"),
        bottom=Side(style="hair", color="CFD8D6"),
    )
    regular_font = Font(name="Roboto", size=10, color="172522")
    centered_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rest_fill = PatternFill("solid", fgColor="E9ECEF")
    vacation_fill = PatternFill("solid", fgColor="FFF0CC")
    rests_fill = PatternFill("solid", fgColor="F4F7F6")

    for row_index, (pid, display_name) in enumerate(people, start=3):
        name_cell = sheet.cell(row=row_index, column=1, value=display_name)
        name_cell.font = Font(name="Roboto", size=10, bold=True, color="172522")
        name_cell.alignment = Alignment(vertical="center", wrap_text=True)
        name_cell.border = grid_border
        rest_count = 0
        vacation = vacations.get(pid)
        if vacation and isinstance(vacation[0], str):
            vacation = [vacation]

        for column, assigned_date in enumerate(days, start=2):
            cell = sheet.cell(row=row_index, column=column)
            cell.font = regular_font
            assignment = assignments_by_day.get((assigned_date, pid))
            if vacation and any(
                start <= assigned_date.isoformat() <= end
                for start, end in vacation
            ):
                cell.value = "Vacaciones"
                cell.fill = vacation_fill
            elif assignment is None:
                cell.value = ""
            else:
                branch_id, start_time, end_time, assignment_type = assignment
                if branch_id is None:
                    cell.value = (
                        "Incapacidad"
                        if assignment_type == "incapacity"
                        else "Descanso"
                    )
                    cell.fill = rest_fill
                    if assignment_type == "rest":
                        rest_count += 1
                else:
                    branch_name, branch_color = branches.get(branch_id, ("Sucursal eliminada", "#E9ECEF"))
                    cell.value = f"{branch_name}\n{start_time}–{end_time}"
                    color = str(branch_color).lstrip("#")
                    if len(color) == 6:
                        cell.fill = PatternFill("solid", fgColor=color.upper())
                        red, green, blue = (int(color[index:index + 2], 16) for index in (0, 2, 4))
                        luminance = 0.299 * red + 0.587 * green + 0.114 * blue
                        cell.font = Font(name="Roboto", size=9, color="172522" if luminance > 155 else "FFFFFF")
            cell.alignment = centered_wrap
            cell.border = grid_border

        rest_cell = sheet.cell(row=row_index, column=last_column, value=rest_count)
        rest_cell.font = regular_font
        rest_cell.fill = rests_fill
        rest_cell.alignment = centered_wrap
        rest_cell.border = grid_border
        sheet.row_dimensions[row_index].height = 38

    sheet.column_dimensions["A"].width = 30
    for column in range(2, last_column):
        sheet.column_dimensions[sheet.cell(row=2, column=column).column_letter].width = 19
    sheet.column_dimensions[sheet.cell(row=2, column=last_column).column_letter].width = 12
    sheet.auto_filter.ref = f"A2:{sheet.cell(row=max(2, len(people) + 2), column=last_column).coordinate}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
