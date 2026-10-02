"""
Gestión de turnos semanales (Flet, desktop) con SQLite.
Personas y sucursales editables · horarios personalizados · descansos · Excel.

Instalar y ejecutar:
    pip install flet==0.28.3 openpyxl
    python turnos_app_actualizado.py

La base de datos (turnos.db) se crea junto al script. Si ya existe una base de
datos de una versión anterior, se migra automáticamente para admitir DNI y
horarios personalizados.
"""
import base64
import calendar
import sqlite3
from datetime import date, timedelta
from io import BytesIO

import flet as ft

from analytics import shift_metrics, summarize_period
from app_config import (
    ANALYSIS_PASSWORD,
    AVATAR_COLORS,
    BRANCH_PALETTE,
    DAY_SHIFT_END_HOUR,
    DAY_SHIFT_START_HOUR,
    DAYS,
    DB_PATH,
    LEGACY_SHIFTS,
    MONTHS,
    PAGE_BG,
    REST_COLOR,
    SEED_BRANCHES,
    SEED_PEOPLE,
    SELECTED_WEEK_KEY,
    WORK_SCHEDULES,
    _pick_color,
    clean_dni,
    monday_of,
    normalize_time,
)
from database import Database
from schedule_export import build_schedule_workbook

# ======================= Interfaz =======================
def main(page: ft.Page):
    db = Database(DB_PATH)
    data = {}

    def reload():
        data["people"] = db.people()
        data["vacations"] = db.vacations()
        data["branches"] = db.branches()
        data["branch"] = {b[0]: (b[1], b[2]) for b in data["branches"]}
        data["branch_id"] = {b[1]: b[0] for b in data["branches"]}
        data["name_counts"] = {}
        for _, name, _ in data["people"]:
            key = name.strip().lower()
            data["name_counts"][key] = data["name_counts"].get(key, 0) + 1

    def stored_week():
        try:
            raw = page.client_storage.get(SELECTED_WEEK_KEY)
            if raw:
                return monday_of(date.fromisoformat(str(raw)))
        except (TypeError, ValueError):
            pass
        except Exception:
            # Si el almacenamiento local no está disponible, la app sigue usando la semana actual.
            pass
        return monday_of(date.today())

    def remember_week():
        try:
            page.client_storage.set(SELECTED_WEEK_KEY, week_key())
        except Exception:
            # La selección permanece durante esta sesión aunque no pueda persistirse.
            pass

    reload()
    today = date.today()
    state = {
        "section": 0,
        "week_start": stored_week(),
        "analysis_month": today.strftime("%Y-%m"),
        "fortnight": 1 if today.day <= 15 else 2,
        "export_period": None,
        "export_bytes": None,
        "export_success_message": "Excel de la quincena guardado correctamente.",
        "analysis_authenticated": False,
    }

    # ---------- tema ----------
    page.title = "Turnos"
    page.padding = 0
    page.spacing = 0
    page.fonts = {"Roboto": "/fonts/Roboto-Regular.ttf"}
    page.theme = ft.Theme(color_scheme_seed="#0F766E", font_family="Roboto", use_material3=True)
    page.dark_theme = ft.Theme(color_scheme_seed="#0F766E", font_family="Roboto", use_material3=True)
    page.theme_mode = ft.ThemeMode.DARK
    page.window.width = 1380
    page.window.height = 880
    page.window.min_width = 1180
    page.window.min_height = 700

    MUTED = ft.Colors.ON_SURFACE_VARIANT
    LINE = ft.Colors.with_opacity(0.55, ft.Colors.OUTLINE_VARIANT)

    def monday():
        return state["week_start"]

    def week_key():
        return monday().isoformat()

    # ---------- helpers de UI ----------
    def toast(msg):
        page.open(ft.SnackBar(ft.Text(msg), behavior=ft.SnackBarBehavior.FLOATING, width=420))

    def confirm(title, text, on_yes, label="Eliminar", danger=True):
        def yes(e):
            page.close(dlg)
            on_yes()

        style = ft.ButtonStyle(bgcolor=ft.Colors.ERROR, color=ft.Colors.ON_ERROR) if danger else None
        dlg = ft.AlertDialog(
            modal=True,
            shape=ft.RoundedRectangleBorder(radius=14),
            title=ft.Text(title),
            content=ft.Container(ft.Text(text, color=MUTED), width=380),
            actions=[ft.TextButton("Cancelar", on_click=lambda e: page.close(dlg)),
                     ft.FilledButton(label, on_click=yes, style=style)],
        )
        page.open(dlg)

    def surface(content, expand=False, padding=0, radius=14):
        return ft.Container(
            content=content, expand=expand, padding=padding, border_radius=radius,
            bgcolor=ft.Colors.SURFACE, border=ft.border.all(1, LINE),
        )

    def hover_scale(e):
        e.control.scale = 1.025 if e.data == "true" else 1
        e.control.update()

    def avatar(pid, name, size=34):
        initials = "".join(w[0] for w in name.split()[:2]).upper() or "?"
        color = AVATAR_COLORS[pid % len(AVATAR_COLORS)]
        return ft.Container(
            width=size, height=size, border_radius=size / 2, alignment=ft.alignment.center,
            bgcolor=ft.Colors.with_opacity(0.16, color),
            content=ft.Text(initials, size=size * 0.38, weight=ft.FontWeight.W_700, color=color),
        )

    def person_label(pid, name, dni):
        if data["name_counts"].get(name.strip().lower(), 0) > 1:
            return f"{name} · DNI {dni}" if dni else f"{name} · sin DNI"
        return name

    def format_date_label(value):
        return date.fromisoformat(value).strftime("%d/%m/%Y") if value else ""

    def is_vacation_day(pid, assigned_date):
        vacation = data["vacations"].get(pid)
        return bool(vacation and vacation[0] <= assigned_date.isoformat() <= vacation[1])

    def week_range_label(start):
        end = start + timedelta(days=6)
        if start.month == end.month:
            return f"{start.day}–{end.day} {MONTHS[end.month - 1]} {end.year}"
        return f"{start.day} {MONTHS[start.month - 1]} – {end.day} {MONTHS[end.month - 1]} {end.year}"

    def week_label(start=None):
        start = start or monday()
        return f"Semana {start.isocalendar().week:02d} · {week_range_label(start)}"

    def page_title(title, subtitle, *actions):
        return ft.Row(
            [ft.Column([ft.Text(title, size=26, weight=ft.FontWeight.W_700),
                        ft.Text(subtitle, size=13, color=MUTED)], spacing=2),
             ft.Container(expand=True), *actions],
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
        )

    def period_dates():
        year, month = map(int, state["analysis_month"].split("-"))
        month_end = calendar.monthrange(year, month)[1]
        if state["fortnight"] == 1:
            return date(year, month, 1), date(year, month, 15)
        return date(year, month, 16), date(year, month, month_end)

    def period_data(start_date, end_date):
        assignments = db.assignments_between(start_date, end_date)
        return assignments, *summarize_period(data["people"], assignments)

    # ---------- exportación a Excel ----------
    def export_fortnight_excel(start_date, end_date):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise RuntimeError("Para descargar el Excel instala openpyxl: pip install openpyxl") from exc

        wb = Workbook()
        ws = wb.active
        ws.title = "Quincena"
        assignments, summaries, totals = period_data(start_date, end_date)
        header_fill = PatternFill("solid", fgColor="0F766E")
        rest_fill = PatternFill("solid", fgColor="EEF1F4")
        total_fill = PatternFill("solid", fgColor="DCEFEA")

        ws.append([f"Resumen de horas · {start_date:%d/%m/%Y} al {end_date:%d/%m/%Y}"])
        ws.merge_cells("A1:F1")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="left")
        ws.append(["Nombre", "Días trabajados", "Trabajadas", "Horas extras", "HD", "HN"])
        for cell in ws[2]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")

        for pid, pname, dni in data["people"]:
            person_metrics = summaries[pid]
            ws.append([
                person_label(pid, pname, dni), person_metrics["days"], person_metrics["worked"],
                person_metrics["overtime"], person_metrics["daytime"], person_metrics["nighttime"],
            ])
        ws.append(["TOTAL"] + [totals[key] for key in
                               ("days", "worked", "overtime", "daytime", "nighttime")])
        for cell in ws[ws.max_row]:
            cell.fill = total_fill
            cell.font = Font(bold=True)

        ws.append([])
        ws.append(["Detalle de turnos"])
        ws[ws.max_row][0].font = Font(bold=True, size=12)
        ws.append(["Nombre", "Día", "Sucursal", "Horario inicio", "Horario fin", "Trabajadas"])
        detail_header = ws.max_row
        for cell in ws[detail_header]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")

        people_by_id = {pid: (name, dni) for pid, name, dni in data["people"]}
        for assigned_date, pid, branch_id, start_time, end_time in assignments:
            pname, dni = people_by_id[pid]
            day_label = f"{DAYS[assigned_date.weekday()]} {assigned_date:%d/%m/%Y}"
            if branch_id is None:
                row = [person_label(pid, pname, dni), day_label, "Descanso", "", "", 0]
            else:
                branch_name = data["branch"].get(branch_id, ("Sucursal eliminada", REST_COLOR))[0]
                worked = shift_metrics(start_time, end_time)[0]
                row = [person_label(pid, pname, dni), day_label, branch_name,
                       start_time, end_time, worked]
            ws.append(row)
            if branch_id is None:
                for cell in ws[ws.max_row]:
                    cell.fill = rest_fill

        for column, width in zip(("A", "B", "C", "D", "E", "F"), (30, 22, 24, 16, 16, 16)):
            ws.column_dimensions[column].width = width
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A{detail_header}:F{ws.max_row}"

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def on_export_result(e: ft.FilePickerResultEvent):
        if not e.path or not state["export_bytes"]:
            return
        path = e.path if e.path.lower().endswith(".xlsx") else f"{e.path}.xlsx"
        try:
            with open(path, "wb") as output:
                output.write(state["export_bytes"])
        except OSError as exc:
            toast(f"No se pudo guardar el Excel: {exc}")
            return
        toast(state["export_success_message"])

    export_picker = ft.FilePicker(on_result=on_export_result)
    page.overlay.append(export_picker)

    def save_export_file(workbook_bytes, file_name, dialog_title, success_message):
        state["export_bytes"] = workbook_bytes
        state["export_success_message"] = success_message
        if page.web:
            try:
                export_picker.save_file(
                    dialog_title=dialog_title,
                    file_name=file_name,
                    allowed_extensions=["xlsx"],
                    src_bytes=workbook_bytes,
                )
            except TypeError as exc:
                if "src_bytes" not in str(exc):
                    raise
                payload = base64.b64encode(workbook_bytes).decode("ascii")
                download_url = (
                    "data:application/vnd.openxmlformats-officedocument."
                    f"spreadsheetml.sheet;base64,{payload}"
                )
                page.launch_url(download_url)
            toast("Descarga de la grilla iniciada.")
        else:
            export_picker.save_file(
                dialog_title=dialog_title,
                file_name=file_name,
                allowed_extensions=["xlsx"],
            )

    # ---------- editor de celda ----------
    def open_cell_editor(pid, pname, dni, day):
        day_date = monday() + timedelta(days=day)
        if is_vacation_day(pid, day_date):
            toast("No se pueden asignar turnos durante las vacaciones.")
            return
        if not data["branches"]:
            toast("Crea una sucursal antes de asignar turnos.")
            return
        sch = db.week(week_key())
        assigned = (pid, day) in sch
        branch_id, start_time, end_time = sch.get((pid, day), (None, None, None))

        if not assigned:
            mode = "Turno"
        elif branch_id is None:
            mode = "Descanso"
        else:
            mode = "Turno"

        dd_mode = ft.Dropdown(
            label="Tipo", value=mode, border_radius=10,
            options=[ft.dropdown.Option(x) for x in ("Turno", "Descanso", "Sin asignar")],
        )
        selected_branch = {
            "id": branch_id if branch_id in data["branch"] else data["branches"][0][0]
        }
        branch_chips = []

        def select_branch(selected_id):
            if dd_mode.value != "Turno":
                return
            selected_branch["id"] = selected_id
            for chip_id, chip in branch_chips:
                chip.selected = chip_id == selected_id
                chip.update()

        for current_branch_id, branch_name, branch_color in data["branches"]:
            chip = ft.Chip(
                label=ft.Text(branch_name),
                leading=ft.Container(
                    width=10, height=10, border_radius=5, bgcolor=branch_color,
                ),
                selected=current_branch_id == selected_branch["id"],
                selected_color=ft.Colors.with_opacity(0.18, branch_color),
                show_checkmark=True,
                disabled=mode != "Turno",
                on_click=lambda e, bid=current_branch_id: select_branch(bid),
                data=current_branch_id,
            )
            branch_chips.append((current_branch_id, chip))

        branch_selection = ft.Column([
            ft.Text("Sucursal", size=13, weight=ft.FontWeight.W_600),
            ft.Row([chip for _, chip in branch_chips], wrap=True, spacing=8, run_spacing=6),
        ], tight=True, spacing=6)
        tf_start = ft.TextField(
            label="Horario inicio", value=start_time or "", hint_text="HH:MM",
            border_radius=10, expand=True, autofocus=not assigned, disabled=mode != "Turno",
        )
        tf_end = ft.TextField(
            label="Horario fin", value=end_time or "", hint_text="HH:MM",
            border_radius=10, expand=True, disabled=mode != "Turno",
        )
        schedule_row = ft.Row([tf_start, tf_end], spacing=10)
        help_text = ft.Text("Puedes usar un horario que termine al día siguiente, por ejemplo 22:00 – 06:00.",
                            size=12, color=MUTED, visible=mode == "Turno")

        def refresh(_=None):
            working = dd_mode.value == "Turno"
            tf_start.disabled = not working
            tf_end.disabled = not working
            for _, chip in branch_chips:
                chip.disabled = not working
            help_text.visible = working
            if _ is not None:
                tf_start.update()
                tf_end.update()
                for _, chip in branch_chips:
                    chip.update()
                help_text.update()

        dd_mode.on_change = refresh

        def clear_errors():
            for control in (tf_start, tf_end):
                control.error_text = None

        def save(e):
            clear_errors()
            if dd_mode.value == "Sin asignar":
                db.clear_assignment(week_key(), pid, day)
            elif dd_mode.value == "Descanso":
                db.set_assignment(week_key(), pid, day, None, None, None)
            else:
                start = normalize_time(tf_start.value)
                end = normalize_time(tf_end.value)
                invalid = False
                if not start:
                    tf_start.error_text = "Usa formato HH:MM."
                    invalid = True
                if not end:
                    tf_end.error_text = "Usa formato HH:MM."
                    invalid = True
                if invalid:
                    dlg.update()
                    return
                if start == end:
                    tf_end.error_text = "Debe ser distinta de la hora de inicio."
                    dlg.update()
                    return
                try:
                    db.set_assignment(week_key(), pid, day, selected_branch["id"], start, end)
                except ValueError as exc:
                    tf_end.error_text = str(exc)
                    dlg.update()
                    return
            page.close(dlg)
            show()

        tf_start.on_submit = save
        tf_end.on_submit = save
        display_name = person_label(pid, pname, dni)
        dlg = ft.AlertDialog(
            modal=True, shape=ft.RoundedRectangleBorder(radius=14),
            title=ft.Row([avatar(pid, pname, 36),
                          ft.Column([ft.Text(display_name, size=17, weight=ft.FontWeight.W_600),
                                     ft.Text(f"{DAYS[day]} {day_date.day} {MONTHS[day_date.month - 1]}",
                                             size=12, color=MUTED)], spacing=0)], spacing=12),
            content=ft.Container(
                ft.Column([dd_mode, schedule_row, branch_selection, help_text], tight=True, spacing=12),
                width=380,
            ),
            actions=[ft.TextButton("Cancelar", on_click=lambda e: page.close(dlg)),
                     ft.FilledButton("Guardar", on_click=save)],
        )
        page.open(dlg)

    # ---------- celdas del horario ----------
    def person_cell(pid, pname, dni, day, sch):
        key = (pid, day)
        assigned_date = monday() + timedelta(days=day)
        if is_vacation_day(pid, assigned_date):
            return ft.Container(
                expand=1, height=58, border_radius=10,
                alignment=ft.alignment.center,
                bgcolor=ft.Colors.with_opacity(0.12, ft.Colors.PRIMARY),
                border=ft.border.all(1, ft.Colors.with_opacity(0.35, ft.Colors.PRIMARY)),
                tooltip=f"Vacaciones · {format_date_label(assigned_date.isoformat())}",
                content=ft.Text("Vacaciones", size=11, weight=ft.FontWeight.W_600,
                                color=ft.Colors.PRIMARY),
            )
        base = dict(
            expand=1, height=58, border_radius=10,
            animate_scale=ft.Animation(110, ft.AnimationCurve.EASE_OUT),
            on_hover=hover_scale,
            on_click=lambda e: open_cell_editor(pid, pname, dni, day),
        )
        if key not in sch:
            return ft.Container(
                **base, tooltip="Asignar turno", alignment=ft.alignment.center,
                border=ft.border.all(1, LINE),
                content=ft.Icon(ft.Icons.ADD, size=16, color=ft.Colors.with_opacity(0.5, MUTED)),
            )

        branch_id, start_time, end_time = sch[key]
        rest = branch_id is None
        branch_name, branch_color = data["branch"].get(branch_id, ("—", REST_COLOR))
        color = REST_COLOR if rest else branch_color
        hours = "Libre" if rest else f"{start_time} – {end_time}"
        tooltip = "Descanso" if rest else f"{branch_name} · {hours}"
        return ft.Container(
            **base, padding=ft.padding.symmetric(vertical=8, horizontal=8),
            bgcolor=ft.Colors.with_opacity(0.10 if rest else 0.15, color),
            tooltip=tooltip,
            content=ft.Row(
                [ft.Container(width=4, border_radius=4, bgcolor=color),
                 ft.Column(
                     [ft.Text("Descanso" if rest else branch_name, size=12, weight=ft.FontWeight.W_600,
                              max_lines=1, overflow=ft.TextOverflow.ELLIPSIS,
                              color=MUTED if rest else None),
                      ft.Row([ft.Icon(ft.Icons.SCHEDULE_OUTLINED, size=12, color=color),
                              ft.Text(hours, size=11, color=MUTED)],
                             spacing=4)],
                     spacing=2, expand=True, alignment=ft.MainAxisAlignment.CENTER)],
                spacing=8, vertical_alignment=ft.CrossAxisAlignment.STRETCH),
        )

    # ---------- sección: Turnos ----------
    def build_schedule():
        sch = db.week(week_key())
        current_monday = monday_of(date.today())
        today_idx = date.today().weekday() if monday() == current_monday else None
        people = data["people"]

        def head_row():
            start = monday()
            cells = [ft.Container(ft.Text("Persona", size=12, color=MUTED, weight=ft.FontWeight.W_600), width=220)]
            for i, d in enumerate(DAYS):
                is_today = i == today_idx
                cells.append(ft.Container(
                    expand=1, alignment=ft.alignment.center, padding=ft.padding.symmetric(vertical=4),
                    border_radius=8,
                    bgcolor=_pick_color("PRIMARY_CONTAINER", "PRIMARY") if is_today else None,
                    content=ft.Text(f"{d} {(start + timedelta(days=i)).day}", size=12,
                                    weight=ft.FontWeight.W_700 if is_today else ft.FontWeight.W_600,
                                    color=_pick_color("ON_PRIMARY_CONTAINER", "ON_PRIMARY") if is_today else MUTED)))
            cells.append(ft.Container(ft.Text("Descansos", size=12, color=MUTED, weight=ft.FontWeight.W_600),
                                      width=76, alignment=ft.alignment.center))
            return ft.Row(cells, spacing=6)

        rows = []
        for pid, pname, dni in people:
            available_days = [
                d for d in range(7)
                if not is_vacation_day(pid, monday() + timedelta(days=d))
            ]
            assigned = [sch[(pid, d)] for d in available_days if (pid, d) in sch]
            rests = sum(1 for branch_id, _, _ in assigned if branch_id is None)
            alert = bool(available_days) and len(assigned) == len(available_days) and rests == 0
            display_name = person_label(pid, pname, dni)
            badge = ft.Container(
                width=76, alignment=ft.alignment.center,
                content=ft.Container(
                    padding=ft.padding.symmetric(horizontal=10, vertical=3), border_radius=12,
                    bgcolor=ft.Colors.with_opacity(0.14, ft.Colors.ERROR if alert else REST_COLOR),
                    tooltip="Sin día de descanso" if alert else None,
                    content=ft.Text(f"{rests}", size=12, weight=ft.FontWeight.W_700,
                                    color=ft.Colors.ERROR if alert else MUTED)))
            rows.append(ft.Row(
                [ft.Container(ft.Row([avatar(pid, pname, 30),
                                      ft.Text(display_name, size=13, expand=True)],
                                     spacing=8), width=220)]
                + [person_cell(pid, pname, dni, d, sch) for d in range(7)] + [badge], spacing=6))

        if not rows:
            rows = [ft.Container(padding=40, alignment=ft.alignment.center, content=ft.Column(
                [ft.Icon(ft.Icons.INBOX_OUTLINED, size=36, color=MUTED),
                 ft.Text("No hay personas todavía.", color=MUTED),
                 ft.Text("Agrégalas en la sección Personas.", size=12, color=MUTED)],
                horizontal_alignment=ft.CrossAxisAlignment.CENTER))]

        pending = sum(
            1 for pid, _, _ in people for d in range(7)
            if not is_vacation_day(pid, monday() + timedelta(days=d)) and (pid, d) not in sch
        )
        no_rest = 0
        for pid, _, _ in people:
            available_days = [
                d for d in range(7)
                if not is_vacation_day(pid, monday() + timedelta(days=d))
            ]
            assigned = [sch[(pid, d)] for d in available_days if (pid, d) in sch]
            if (available_days and len(assigned) == len(available_days)
                    and all(branch_id is not None for branch_id, _, _ in assigned)):
                no_rest += 1
        status = ft.Row([
            ft.Row([ft.Icon(ft.Icons.PEOPLE_OUTLINE, size=16, color=MUTED),
                    ft.Text(f"{len(people)} personas", size=12, color=MUTED)], spacing=6),
            ft.Row([ft.Icon(ft.Icons.WARNING_AMBER_ROUNDED, size=16,
                            color=ft.Colors.ERROR if no_rest else MUTED),
                    ft.Text(f"{no_rest} sin descanso", size=12,
                            color=ft.Colors.ERROR if no_rest else MUTED)], spacing=6),
            ft.Row([ft.Icon(ft.Icons.EVENT_BUSY_OUTLINED, size=16, color=MUTED),
                    ft.Text(f"{pending} celdas sin asignar", size=12, color=MUTED)], spacing=6),
        ], spacing=22)

        # --- acciones de semana ---
        def do_copy():
            db.copy_week((monday() - timedelta(weeks=1)).isoformat(), week_key())
            show()
            toast("Semana copiada desde la anterior.")

        def do_auto():
            db.autogenerate(week_key())
            show()
            toast("Semana generada con datos de ejemplo.")

        def do_clear():
            db.clear_week(week_key())
            show()
            toast("Semana vaciada.")

        menu = ft.PopupMenuButton(
            icon=ft.Icons.MORE_VERT, tooltip="Acciones de la semana",
            items=[
                ft.PopupMenuItem(text="Copiar semana anterior", icon=ft.Icons.CONTENT_COPY_OUTLINED,
                                 on_click=lambda e: confirm("Copiar semana anterior",
                                                            "Se reemplazarán los turnos de esta semana.",
                                                            do_copy, "Copiar", danger=False)),
                ft.PopupMenuItem(text="Generar datos de ejemplo", icon=ft.Icons.AUTO_AWESOME_OUTLINED,
                                 on_click=lambda e: confirm("Generar datos de ejemplo",
                                                            "Se reemplazarán los turnos de esta semana.",
                                                            do_auto, "Generar", danger=False)),
                ft.PopupMenuItem(text="Vaciar semana", icon=ft.Icons.DELETE_SWEEP_OUTLINED,
                                 on_click=lambda e: confirm("Vaciar semana",
                                                            "Se quitarán todos los turnos de esta semana.",
                                                            do_clear, "Vaciar")),
            ])

        def go(delta):
            state["week_start"] = monday() + timedelta(weeks=delta)
            remember_week()
            show()

        def go_today(e):
            state["week_start"] = current_monday
            remember_week()
            show()

        def week_options():
            options = {current_monday + timedelta(weeks=i) for i in range(-12, 27)}
            for value in db.available_weeks():
                try:
                    options.add(monday_of(date.fromisoformat(value)))
                except ValueError:
                    continue
            options.add(monday())
            return sorted(options)

        def select_week(e):
            if not e.control.value:
                return
            try:
                state["week_start"] = monday_of(date.fromisoformat(e.control.value))
            except ValueError:
                return
            remember_week()
            show()

        def open_schedule_export(e):
            selected_dates = {"start": monday(), "end": monday() + timedelta(days=6)}
            start_field = ft.TextField(
                label="Desde", value=format_date_label(selected_dates["start"].isoformat()),
                read_only=True, expand=True, border_radius=10,
            )
            end_field = ft.TextField(
                label="Hasta", value=format_date_label(selected_dates["end"].isoformat()),
                read_only=True, expand=True, border_radius=10,
            )
            range_error = ft.Text("", size=12, color=ft.Colors.ERROR)

            def choose_date(key):
                def selected(event):
                    value = event.control.value
                    if value is None:
                        return
                    selected_dates[key] = value.date() if hasattr(value, "date") else value
                    field = start_field if key == "start" else end_field
                    field.value = format_date_label(selected_dates[key].isoformat())
                    range_error.value = ""
                    dialog.update()

                picker = ft.DatePicker(
                    value=selected_dates[key],
                    first_date=date(1900, 1, 1),
                    last_date=date(2100, 12, 31),
                    help_text="Selecciona la fecha",
                    confirm_text="Aceptar",
                    cancel_text="Cancelar",
                    on_change=selected,
                )
                page.open(picker)

            def export_range(event):
                start_date = selected_dates["start"]
                end_date = selected_dates["end"]
                if end_date < start_date:
                    range_error.value = "La fecha final debe ser igual o posterior a la inicial."
                    dialog.update()
                    return
                people_for_export = [
                    (pid, person_label(pid, name, dni)) for pid, name, dni in data["people"]
                ]
                try:
                    workbook_bytes = build_schedule_workbook(
                        start_date,
                        end_date,
                        people_for_export,
                        db.assignments_between(start_date, end_date),
                        data["branch"],
                        data["vacations"],
                    )
                except (RuntimeError, ValueError) as exc:
                    range_error.value = str(exc)
                    dialog.update()
                    return

                file_name = f"grilla_turnos_{start_date:%Y%m%d}_{end_date:%Y%m%d}.xlsx"
                page.close(dialog)
                try:
                    save_export_file(
                        workbook_bytes,
                        file_name,
                        "Guardar grilla de turnos",
                        "Grilla de turnos guardada correctamente.",
                    )
                except Exception as exc:
                    toast(f"No se pudo guardar el Excel: {exc}")

            dialog = ft.AlertDialog(
                modal=True,
                shape=ft.RoundedRectangleBorder(radius=14),
                title=ft.Text("Descargar grilla de turnos"),
                content=ft.Container(
                    ft.Column([
                        ft.Text("Selecciona el rango que quieres incluir en el Excel.", size=13, color=MUTED),
                        ft.Row([
                            start_field,
                            ft.IconButton(ft.Icons.CALENDAR_MONTH, tooltip="Elegir fecha inicial",
                                          on_click=lambda event: choose_date("start")),
                        ], spacing=4),
                        ft.Row([
                            end_field,
                            ft.IconButton(ft.Icons.CALENDAR_MONTH, tooltip="Elegir fecha final",
                                          on_click=lambda event: choose_date("end")),
                        ], spacing=4),
                        range_error,
                    ], tight=True, spacing=8),
                    width=400,
                ),
                actions=[
                    ft.TextButton("Cancelar", on_click=lambda event: page.close(dialog)),
                    ft.FilledButton("Descargar", icon=ft.Icons.DOWNLOAD_OUTLINED, on_click=export_range),
                ],
            )
            page.open(dialog)

        def week_option(week_start):
            week_end = week_start + timedelta(days=6)
            return ft.dropdown.Option(
                key=week_start.isoformat(),
                text=f"Sem {week_start.isocalendar().week:02d} · "
                     f"{week_start.day} {MONTHS[week_start.month - 1]} – "
                     f"{week_end.day} {MONTHS[week_end.month - 1]}",
            )

        week_select = ft.Dropdown(
            value=week_key(),
            options=[week_option(w) for w in week_options()],
            on_change=select_week,
            width=240,
            border_radius=8,
        )
        week_nav = ft.Container(
            border_radius=22, bgcolor=ft.Colors.SURFACE, border=ft.border.all(1, LINE),
            content=ft.Row([
                ft.IconButton(ft.Icons.CHEVRON_LEFT, icon_size=20, tooltip="Semana anterior",
                              on_click=lambda e: go(-1)),
                week_select,
                ft.IconButton(ft.Icons.CHEVRON_RIGHT, icon_size=20, tooltip="Semana siguiente",
                              on_click=lambda e: go(1)),
            ], spacing=0))

        top = page_title(
            "Turnos", "Haz clic en una celda para asignar sucursal y horario.",
            ft.TextButton("Hoy", on_click=go_today, visible=monday() != current_monday),
            week_nav,
            ft.OutlinedButton("Descargar grilla", icon=ft.Icons.DOWNLOAD_OUTLINED,
                              on_click=open_schedule_export),
            menu)

        grid = surface(
            ft.Column([
                head_row(),
                ft.Divider(height=1, color=LINE),
                ft.Column(rows, scroll=ft.ScrollMode.AUTO, expand=True, spacing=5),
            ], spacing=6, expand=True),
            expand=True, padding=ft.padding.only(left=12, right=12, top=8, bottom=10))

        return ft.Column([top, status, grid], expand=True, spacing=8)

    # ---------- sección: Análisis ----------
    def build_analysis():
        month_values = set()
        current_month_index = today.year * 12 + today.month - 1
        for offset in range(-24, 25):
            year, month_index = divmod(current_month_index + offset, 12)
            month_values.add(f"{year:04d}-{month_index + 1:02d}")
        for week_start in db.available_weeks():
            try:
                week_date = date.fromisoformat(week_start)
            except ValueError:
                continue
            for day_offset in range(7):
                assigned_date = week_date + timedelta(days=day_offset)
                month_values.add(assigned_date.strftime("%Y-%m"))

        month_options = []
        for value in sorted(month_values):
            year, month = map(int, value.split("-"))
            month_options.append(ft.dropdown.Option(
                key=value, text=f"{MONTHS[month - 1].title()} {year}"
            ))

        def change_month(e):
            if e.control.value:
                state["analysis_month"] = e.control.value
                show()

        month_dropdown = ft.Dropdown(
            label="Mes", value=state["analysis_month"], options=month_options,
            on_change=change_month, width=180, border_radius=12,
        )
        year, month = map(int, state["analysis_month"].split("-"))
        month_end = calendar.monthrange(year, month)[1]

        def change_fortnight(e):
            if e.control.value:
                state["fortnight"] = int(e.control.value)
                show()

        fortnight_dropdown = ft.Dropdown(
            label="Quincena", value=str(state["fortnight"]),
            options=[ft.dropdown.Option(key="1", text="1 al 15"),
                     ft.dropdown.Option(key="2", text=f"16 al {month_end}")],
            on_change=change_fortnight, width=170, border_radius=12,
        )

        start_date, end_date = period_dates()
        assignments, summaries, totals = period_data(start_date, end_date)

        def choose_export_path(e):
            state["export_period"] = (start_date, end_date)
            file_name = f"turnos_{start_date:%Y-%m}_{start_date.day:02d}-{end_date.day:02d}.xlsx"
            try:
                workbook_bytes = export_fortnight_excel(start_date, end_date)
                if page.web:
                    try:
                        export_picker.save_file(
                            dialog_title="Descargar turnos de la quincena",
                            file_name=file_name,
                            allowed_extensions=["xlsx"],
                            src_bytes=workbook_bytes,
                        )
                    except TypeError as exc:
                        if "src_bytes" not in str(exc):
                            raise
                        payload = base64.b64encode(workbook_bytes).decode("ascii")
                        download_url = (
                            "data:application/vnd.openxmlformats-officedocument."
                            f"spreadsheetml.sheet;base64,{payload}"
                        )
                        page.launch_url(download_url)
                    toast("Descarga de Excel iniciada.")
                else:
                    state["export_bytes"] = workbook_bytes
                    export_picker.save_file(
                        dialog_title="Guardar turnos de la quincena",
                        file_name=file_name,
                        allowed_extensions=["xlsx"],
                    )
            except Exception as exc:
                toast(str(exc))
                return

        export_button = ft.FilledButton(
            "Descargar Excel", icon=ft.Icons.DOWNLOAD_OUTLINED, on_click=choose_export_path
        )

        def hours_cell(value, bold=False):
            return ft.DataCell(ft.Text(f"{value:.2f}", weight=ft.FontWeight.W_700 if bold else None))

        rows = []
        for pid, pname, dni in data["people"]:
            person_metrics = summaries[pid]
            rows.append(ft.DataRow(cells=[
                ft.DataCell(ft.Text(person_label(pid, pname, dni))),
                ft.DataCell(ft.Text(str(person_metrics["days"]))),
                hours_cell(person_metrics["worked"]),
                hours_cell(person_metrics["overtime"]),
                hours_cell(person_metrics["daytime"]),
                hours_cell(person_metrics["nighttime"]),
            ]))
        rows.append(ft.DataRow(cells=[
            ft.DataCell(ft.Text("TOTAL", weight=ft.FontWeight.W_700)),
            ft.DataCell(ft.Text(str(totals["days"]), weight=ft.FontWeight.W_700)),
            hours_cell(totals["worked"], bold=True),
            hours_cell(totals["overtime"], bold=True),
            hours_cell(totals["daytime"], bold=True),
            hours_cell(totals["nighttime"], bold=True),
        ]))

        table = ft.DataTable(
            columns=[ft.DataColumn(ft.Text(label)) for label in
                     ("Nombre", "Días trabajados", "Trabajadas", "Horas extras", "HD", "HN")],
            rows=rows,
            column_spacing=36,
        )
        subtitle = (
            f"{start_date:%d/%m/%Y} al {end_date:%d/%m/%Y} · Extras sobre 8 h por turno · "
            f"HD {DAY_SHIFT_START_HOUR:02d}:00–{DAY_SHIFT_END_HOUR:02d}:00 · "
            f"HN {DAY_SHIFT_END_HOUR:02d}:00–{DAY_SHIFT_START_HOUR:02d}:00"
        )
        top = page_title("Análisis", subtitle, month_dropdown, fortnight_dropdown, export_button)
        return ft.Column([
            top,
            surface(ft.Column([table], scroll=ft.ScrollMode.AUTO, expand=True),
                    expand=True, padding=ft.padding.symmetric(horizontal=8, vertical=6)),
        ], expand=True, spacing=16)

    # ---------- sección: Personas ----------
    def person_dialog(pid=None, current="", current_dni="", current_vacation=None):
        vacation_start, vacation_end = current_vacation or (None, None)
        vacation_dates = {"start": vacation_start, "end": vacation_end}

        tf_name = ft.TextField(label="Nombre completo", value=current, autofocus=True, border_radius=10)
        tf_dni = ft.TextField(label="DNI (opcional)", value=current_dni, border_radius=10)
        tf_vacation_start = ft.TextField(
            label="Desde", value=format_date_label(vacation_start), read_only=True, expand=True, border_radius=10
        )
        tf_vacation_end = ft.TextField(
            label="Hasta", value=format_date_label(vacation_end), read_only=True, expand=True, border_radius=10
        )
        vacation_error = ft.Text("", size=12, color=ft.Colors.ERROR)

        def choose_vacation_date(key):
            def selected(e):
                value = e.control.value
                if value is None:
                    return
                vacation_dates[key] = value.strftime("%Y-%m-%d")
                field = tf_vacation_start if key == "start" else tf_vacation_end
                field.value = format_date_label(vacation_dates[key])
                vacation_error.value = ""
                dlg.update()

            picker = ft.DatePicker(
                value=date.fromisoformat(vacation_dates[key]) if vacation_dates[key] else date.today(),
                first_date=date(1900, 1, 1), last_date=date(2100, 12, 31),
                help_text="Selecciona la fecha", confirm_text="Aceptar", cancel_text="Cancelar",
                on_change=selected,
            )
            page.open(picker)

        def clear_vacation(e):
            vacation_dates["start"] = None
            vacation_dates["end"] = None
            tf_vacation_start.value = ""
            tf_vacation_end.value = ""
            vacation_error.value = ""
            dlg.update()

        vacation_fields = ft.Column([
            ft.Text("Vacaciones", size=13, weight=ft.FontWeight.W_600),
            ft.Row([
                tf_vacation_start,
                ft.IconButton(ft.Icons.CALENDAR_MONTH, tooltip="Elegir fecha inicial",
                              on_click=lambda e: choose_vacation_date("start")),
            ], spacing=4),
            ft.Row([
                tf_vacation_end,
                ft.IconButton(ft.Icons.CALENDAR_MONTH, tooltip="Elegir fecha final",
                              on_click=lambda e: choose_vacation_date("end")),
            ], spacing=4),
            ft.TextButton("Quitar período de vacaciones", on_click=clear_vacation),
            vacation_error,
        ], tight=True, spacing=4)

        def save(e):
            name = " ".join((tf_name.value or "").split())
            dni = clean_dni(tf_dni.value)
            tf_name.error_text = None
            tf_dni.error_text = None
            vacation_error.value = ""
            if not name:
                tf_name.error_text = "Escribe un nombre."
                dlg.update()
                return
            start = vacation_dates["start"]
            end = vacation_dates["end"]
            if bool(start) != bool(end):
                vacation_error.value = "Selecciona ambas fechas o quita el período."
                dlg.update()
                return
            if start and date.fromisoformat(start) > date.fromisoformat(end):
                vacation_error.value = "La fecha inicial debe ser anterior o igual a la final."
                dlg.update()
                return
            removed_assignments = 0
            try:
                if pid is None:
                    db.add_person(name, dni, start, end)
                else:
                    removed_assignments = db.rename_person(pid, name, dni, start, end)
            except ValueError as exc:
                tf_dni.error_text = str(exc)
                dlg.update()
                return
            except sqlite3.IntegrityError:
                tf_dni.error_text = "Ya existe una persona con ese DNI."
                dlg.update()
                return
            page.close(dlg)
            reload()
            show()
            message = "Persona agregada." if pid is None else "Cambios guardados."
            if removed_assignments:
                message += f" Se quitaron {removed_assignments} asignaciones durante las vacaciones."
            toast(message)

        tf_name.on_submit = save
        tf_dni.on_submit = save
        dlg = ft.AlertDialog(
            modal=True, shape=ft.RoundedRectangleBorder(radius=14),
            title=ft.Text("Nueva persona" if pid is None else "Editar persona"),
            content=ft.Container(ft.Column([tf_name, tf_dni, vacation_fields], tight=True, spacing=10), width=380),
            actions=[ft.TextButton("Cancelar", on_click=lambda e: page.close(dlg)),
                     ft.FilledButton("Agregar" if pid is None else "Guardar cambios", on_click=save)])
        page.open(dlg)

    def delete_person(pid, name):
        def go():
            db.delete_person(pid)
            reload()
            show()
            toast(f"{name} eliminada.")
        confirm(f"Eliminar a {name}", "También se borrarán todos sus turnos de todas las semanas.", go)

    def build_people():
        sch = db.week(week_key())
        items = []
        for pid, name, dni in data["people"]:
            assigned = [sch[(pid, d)] for d in range(7) if (pid, d) in sch]
            rests = sum(1 for branch_id, _, _ in assigned if branch_id is None)
            worked = len(assigned) - rests
            weekly = f"Esta semana: {worked} turnos, {rests} descansos" if assigned else "Sin turnos esta semana"
            identity = f"DNI: {dni}" if dni else "DNI: no registrado"
            vacation = data["vacations"].get(pid)
            details = [ft.Text(f"{identity} · {weekly}", size=12, color=MUTED)]
            if vacation:
                vacation_label = (
                    f"Vacaciones: {format_date_label(vacation[0])} al {format_date_label(vacation[1])}"
                )
                details.append(ft.Text(vacation_label, size=12, color=ft.Colors.PRIMARY))
            display_name = person_label(pid, name, dni)
            items.append(ft.Container(
                padding=ft.padding.symmetric(horizontal=14, vertical=10), border_radius=10,
                content=ft.Row([
                    avatar(pid, name, 38),
                    ft.Column([ft.Text(display_name, size=14, weight=ft.FontWeight.W_600), *details],
                              spacing=0, expand=True),
                    ft.IconButton(ft.Icons.EDIT_OUTLINED, tooltip="Editar",
                                  on_click=lambda e, p=pid, n=name, d=dni, v=vacation: person_dialog(p, n, d, v)),
                    ft.IconButton(ft.Icons.DELETE_OUTLINE, tooltip="Eliminar", icon_color=ft.Colors.ERROR,
                                  on_click=lambda e, p=pid, n=display_name: delete_person(p, n)),
                ], vertical_alignment=ft.CrossAxisAlignment.CENTER)))
            items.append(ft.Divider(height=1, color=LINE))
        if items:
            items.pop()
        else:
            items = [ft.Container(padding=40, alignment=ft.alignment.center,
                                  content=ft.Text("Aún no hay personas. Agrega la primera.", color=MUTED))]
        return ft.Column([
            page_title("Personas", f"{len(data['people'])} en el equipo · DNI opcional",
                       ft.FilledButton("Nueva persona", icon=ft.Icons.ADD, on_click=lambda e: person_dialog())),
            surface(ft.Column(items, scroll=ft.ScrollMode.AUTO, spacing=0), expand=True, padding=6),
        ], expand=True, spacing=16)

    # ---------- sección: Sucursales ----------
    def branch_dialog(bid=None, name="", color=None):
        used = {b[2] for b in data["branches"]}
        sel = {"c": color or next((c for c in BRANCH_PALETTE if c not in used), BRANCH_PALETTE[0])}
        tf = ft.TextField(label="Nombre de la sucursal", value=name, autofocus=True, border_radius=10)
        swatches = ft.Row(wrap=True, spacing=10, run_spacing=10)

        def paint():
            swatches.controls = [
                ft.Container(
                    width=32, height=32, border_radius=16, bgcolor=c,
                    border=ft.border.all(3, ft.Colors.ON_SURFACE if c == sel["c"] else ft.Colors.TRANSPARENT),
                    on_click=lambda e, c=c: pick(c))
                for c in BRANCH_PALETTE]

        def pick(c):
            sel["c"] = c
            paint()
            swatches.update()

        paint()

        def save(e):
            n = " ".join((tf.value or "").split())
            if not n:
                tf.error_text = "Escribe un nombre."
                tf.update()
                return
            try:
                db.add_branch(n, sel["c"]) if bid is None else db.update_branch(bid, n, sel["c"])
            except sqlite3.IntegrityError:
                tf.error_text = "Ya existe una sucursal con ese nombre."
                tf.update()
                return
            page.close(dlg)
            reload()
            show()
            toast("Sucursal agregada." if bid is None else "Cambios guardados.")

        tf.on_submit = save
        dlg = ft.AlertDialog(
            modal=True, shape=ft.RoundedRectangleBorder(radius=14),
            title=ft.Text("Nueva sucursal" if bid is None else "Editar sucursal"),
            content=ft.Container(ft.Column([tf, ft.Text("Color en el horario", size=12, color=MUTED), swatches],
                                           tight=True, spacing=14), width=360),
            actions=[ft.TextButton("Cancelar", on_click=lambda e: page.close(dlg)),
                     ft.FilledButton("Agregar" if bid is None else "Guardar cambios", on_click=save)])
        page.open(dlg)

    def delete_branch(bid, name):
        def go():
            db.delete_branch(bid)
            reload()
            show()
            toast(f"{name} eliminada.")
        confirm(f"Eliminar {name}", "Se borrarán los turnos asignados a esta sucursal en todas las semanas.", go)

    def build_branches():
        sch = db.week(week_key())
        cards = []
        for bid, name, color in data["branches"]:
            mine = [(p, d) for (p, d), (b, _, _) in sch.items() if b == bid]
            people_n = len({p for p, _ in mine})
            cards.append(surface(
                ft.Column([
                    ft.Row([
                        ft.Container(width=40, height=40, border_radius=10, alignment=ft.alignment.center,
                                     bgcolor=ft.Colors.with_opacity(0.16, color),
                                     content=ft.Icon(ft.Icons.STORE_OUTLINED, color=color, size=22)),
                        ft.Column([ft.Text(name, size=15, weight=ft.FontWeight.W_600, max_lines=1,
                                           overflow=ft.TextOverflow.ELLIPSIS),
                                   ft.Text(f"{len(mine)} turnos · {people_n} personas", size=12, color=MUTED)],
                                  spacing=0, expand=True)], spacing=12),
                    ft.Row([ft.Container(expand=True),
                            ft.IconButton(ft.Icons.EDIT_OUTLINED, tooltip="Editar", icon_size=19,
                                          on_click=lambda e, b=bid, n=name, c=color: branch_dialog(b, n, c)),
                            ft.IconButton(ft.Icons.DELETE_OUTLINE, tooltip="Eliminar", icon_size=19,
                                          icon_color=ft.Colors.ERROR,
                                          on_click=lambda e, b=bid, n=name: delete_branch(b, n))],
                           spacing=0),
                ], spacing=4),
                padding=ft.padding.only(left=16, right=8, top=16, bottom=6)))
        for c in cards:
            c.width = 290
        body = ft.Row(cards, wrap=True, spacing=14, run_spacing=14, alignment=ft.MainAxisAlignment.START) if cards \
            else ft.Container(padding=40, content=ft.Text("Aún no hay sucursales. Agrega la primera.", color=MUTED))
        return ft.Column([
            page_title("Sucursales", f"{len(data['branches'])} sucursales · conteo de la {week_label()}",
                       ft.FilledButton("Nueva sucursal", icon=ft.Icons.ADD, on_click=lambda e: branch_dialog())),
            ft.Column([body], scroll=ft.ScrollMode.AUTO, expand=True),
        ], expand=True, spacing=16)

    # ---------- navegación y shell ----------
    main_area = ft.Container(expand=True, padding=ft.padding.only(left=28, right=28, top=22, bottom=22),
                             bgcolor=PAGE_BG)

    def show():
        builders = [build_schedule, build_people, build_branches, build_analysis]
        main_area.content = builders[state["section"]]()
        page.update()

    def open_analysis_login():
        password_field = ft.TextField(
            label="Contraseña", password=True, can_reveal_password=True,
            autofocus=True, border_radius=10,
        )

        def submit(e):
            if password_field.value == ANALYSIS_PASSWORD:
                state["analysis_authenticated"] = True
                state["section"] = 3
                rail.selected_index = 3
                page.close(dialog)
                show()
                return
            password_field.error_text = "Contraseña incorrecta."
            password_field.value = ""
            dialog.update()

        password_field.on_submit = submit
        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("Acceso a Análisis"),
            content=ft.Container(password_field, width=320),
            actions=[
                ft.TextButton(
                    "Cancelar",
                    on_click=lambda e: page.close(dialog),
                ),
                ft.FilledButton("Ingresar", on_click=submit),
            ],
        )
        page.open(dialog)

    def on_nav(e):
        selected_index = e.control.selected_index
        if selected_index == 3 and not state["analysis_authenticated"]:
            rail.selected_index = state["section"]
            page.update()
            open_analysis_login()
            return
        state["section"] = selected_index
        show()

    def toggle_theme(e):
        dark = page.theme_mode == ft.ThemeMode.LIGHT
        page.theme_mode = ft.ThemeMode.DARK if dark else ft.ThemeMode.LIGHT
        e.control.icon = ft.Icons.LIGHT_MODE_OUTLINED if dark else ft.Icons.DARK_MODE_OUTLINED
        page.update()

    rail = ft.NavigationRail(
        selected_index=0, min_width=84, bgcolor=ft.Colors.SURFACE,
        label_type=ft.NavigationRailLabelType.ALL, on_change=on_nav,
        leading=ft.Container(padding=ft.padding.only(top=16, bottom=14), content=ft.Container(
            width=42, height=42, border_radius=12, alignment=ft.alignment.center,
            bgcolor=ft.Colors.PRIMARY, content=ft.Icon(ft.Icons.CALENDAR_VIEW_WEEK, color=ft.Colors.ON_PRIMARY))),
        trailing=ft.Container(padding=ft.padding.only(top=24), content=ft.IconButton(
            ft.Icons.DARK_MODE_OUTLINED, tooltip="Cambiar tema", on_click=toggle_theme)),
        destinations=[
            ft.NavigationRailDestination(icon=ft.Icons.CALENDAR_MONTH_OUTLINED,
                                         selected_icon=ft.Icons.CALENDAR_MONTH, label="Turnos"),
            ft.NavigationRailDestination(icon=ft.Icons.GROUPS_OUTLINED,
                                         selected_icon=ft.Icons.GROUPS, label="Personas"),
            ft.NavigationRailDestination(icon=ft.Icons.STORE_OUTLINED,
                                         selected_icon=ft.Icons.STORE, label="Sucursales"),
            ft.NavigationRailDestination(icon=ft.Icons.INSERT_CHART_OUTLINED,
                                         selected_icon=ft.Icons.INSERT_CHART, label="Análisis"),
        ],
    )

    page.add(ft.Row([rail, ft.VerticalDivider(width=1, color=LINE), main_area], expand=True, spacing=0))
    show()


if __name__ == "__main__":
    ft.app(target=main)
