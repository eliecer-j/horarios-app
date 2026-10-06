from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .metrics import assignment_period_metrics, summarize_period
from .models import Branch
from .services import assignments_between, name_counts, person_display_name


def fortnight_workbook(start_date, end_date, people):
    assignments = assignments_between(start_date, end_date)
    daily_metrics = assignment_period_metrics(people, assignments)
    summaries, totals = summarize_period(people, assignments, daily_metrics)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quincena"
    header_fill = PatternFill("solid", fgColor="0F766E")
    total_fill = PatternFill("solid", fgColor="D9EAE7")
    rest_fill = PatternFill("solid", fgColor="EEF1F4")
    sheet.append([f"Análisis de turnos · {start_date:%d/%m/%Y} al {end_date:%d/%m/%Y}"])
    sheet.merge_cells("A1:F1")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["Nombre", "Días trabajados", "Horas trabajadas", "Horas extra", "Horas diurnas", "Horas nocturnas"])
    for cell in sheet[2]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for person in people:
        metric = summaries[person.pk]
        sheet.append([
            person.name, metric["days"], metric["worked"], metric["overtime"],
            metric["daytime"], metric["nighttime"],
        ])
    sheet.append([
        "TOTAL", totals["days"], totals["worked"], totals["overtime"],
        totals["daytime"], totals["nighttime"],
    ])
    for cell in sheet[sheet.max_row]:
        cell.fill = total_fill
        cell.font = Font(bold=True)
    sheet.append([])
    sheet.append(["Detalle de turnos"])
    sheet[sheet.max_row][0].font = Font(bold=True, size=12)
    sheet.append([
        "Nombre", "Día", "Sucursal", "Horario inicio", "Horario fin", "Trabajadas",
        "Horas extra", "HD", "HN", "Observación",
    ])
    detail_header = sheet.max_row
    for cell in sheet[detail_header]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    people_by_id = {person.pk: person for person in people}
    duplicate_names = name_counts(people)
    branches = {branch.pk: branch.name for branch in Branch.objects.all()}
    day_names = ("Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom")
    for assignment_index, assignment in enumerate(assignments):
        assigned_date, person_id, branch_id, start_time, end_time = assignment[:5]
        actual_start_time = assignment[5] if len(assignment) > 5 else None
        novelty_kind = assignment[6] if len(assignment) > 6 else None
        observation = assignment[7] if len(assignment) > 7 else ""
        metrics = daily_metrics.get(assignment_index, {
            "overtime": 0, "daytime": 0, "nighttime": 0,
        })
        person = people_by_id.get(person_id)
        if person is None:
            continue
        if branch_id is None:
            row = [
                person_display_name(person, duplicate_names),
                f"{day_names[assigned_date.weekday()]} {assigned_date:%d/%m/%Y}",
                "Descanso", "", "", 0, 0, 0, 0, "",
            ]
        else:
            effective_start_time = actual_start_time or start_time
            if novelty_kind in ("did_not_attend", "calamity"):
                hours = 0
                if novelty_kind == "did_not_attend":
                    novelty_label = "No se presentó"
                else:
                    novelty_label = "Calamidad"
                branch_name = f"{branches.get(branch_id, 'Sucursal eliminada')} · {novelty_label}"
            else:
                hours = (int(end_time[:2]) * 60 + int(end_time[3:])) - (
                    int(effective_start_time[:2]) * 60 + int(effective_start_time[3:])
                )
                if hours <= 0:
                    hours += 24 * 60
                branch_name = branches.get(branch_id, "Sucursal eliminada")
            row = [
                person_display_name(person, duplicate_names),
                f"{day_names[assigned_date.weekday()]} {assigned_date:%d/%m/%Y}",
                branch_name,
                effective_start_time,
                end_time,
                hours / 60,
                metrics["overtime"] / 60,
                metrics["daytime"] / 60,
                metrics["nighttime"] / 60,
                observation,
            ]
        sheet.append(row)
        if branch_id is None:
            for cell in sheet[sheet.max_row]:
                cell.fill = rest_fill
    for column in range(1, 11):
        sheet.cell(row=detail_header, column=column).alignment = Alignment(horizontal="center")
    for column, width in zip("ABCDEFGHIJ", (30, 18, 30, 16, 16, 14, 14, 10, 10, 36)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A{detail_header}:J{sheet.max_row}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def workbook_response(payload, filename):
    response = HttpResponse(
        payload,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
