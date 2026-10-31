from datetime import date, timedelta
from io import BytesIO

from django.db import connection
from django.test import TestCase, override_settings
from django.utils import timezone
from django.urls import reverse
from openpyxl import load_workbook

from .exports import fortnight_workbook
from .forms import AssignmentForm, LateArrivalForm
from .metrics import (
    assignment_period_metrics,
    shift_metrics,
    shift_segments_metrics,
    shift_touches_lunch,
    summarize_period,
)
from .models import AssignmentNovelty, AuditLog, Branch, Person, VacationPeriod
from .services import assignments_between, clear_week, delete_branch, monday_of, week_assignments
from .templatetags.schedule_tags import capitalize, signed_class


class LegacyDatabaseViewsTests(TestCase):
    def setUp(self):
        VacationPeriod.objects.all().delete()
        with connection.cursor() as cursor:
            cursor.execute("DELETE FROM assignments")
            cursor.execute("DELETE FROM branches")
            cursor.execute("DELETE FROM people")
        self.person = Person.objects.create(name="Ana Torres")
        self.branch = Branch.objects.create(name="Centro", color="#2F6FED")
        self.week_start = monday_of(date.today())

    def test_schedule_page_renders_legacy_data(self):
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
        response = self.client.get(reverse("schedule", args=[self.week_start.isoformat()]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ana Torres")
        self.assertContains(response, "Centro")
        self.assertContains(response, "schedule-grid-scroll")
        self.assertContains(response, 'title="Ana Torres"', html=False)
        self.assertLess(
            response.content.index(b">Dom<span>"),
            response.content.index(b'class="rest-count-column">Descansos</th>'),
        )

    def test_schedule_filters_rows_by_branch(self):
        other_person = Person.objects.create(name="Beto Ruiz")
        other_branch = Branch.objects.create(name="Norte", color="#EB5757")
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), other_person.pk, other_branch.pk],
            )

        filtered = self.client.get(
            reverse("schedule", args=[self.week_start.isoformat()]),
            {"branch": str(self.branch.pk)},
        )
        self.assertEqual(filtered.status_code, 200)
        self.assertContains(filtered, "Ana Torres")
        self.assertNotContains(filtered, "Beto Ruiz")
        self.assertContains(filtered, f'value="{self.branch.pk}" selected')
        self.assertContains(filtered, "Todas las sucursales")

        unfiltered = self.client.get(
            reverse("schedule", args=[self.week_start.isoformat()]),
        )
        self.assertContains(unfiltered, "Ana Torres")
        self.assertContains(unfiltered, "Beto Ruiz")

    def test_saving_shift_does_not_auto_filter_by_its_branch(self):
        other_person = Person.objects.create(name="Beto Ruiz")
        other_branch = Branch.objects.create(name="Norte", color="#EB5757")
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), other_person.pk, other_branch.pk],
            )

        response = self.client.post(
            reverse("assignment_save", args=[self.week_start.isoformat(), self.person.pk, 0]),
            {
                "status": "work",
                "branch": self.branch.pk,
                "start_time": "08:00",
                "end_time": "16:00",
                "branch_filter": "",
            },
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ana Torres")
        self.assertContains(response, "Beto Ruiz")

    def test_saving_shift_keeps_the_selected_branch_filter(self):
        other_person = Person.objects.create(name="Beto Ruiz")
        other_branch = Branch.objects.create(name="Norte", color="#EB5757")
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), other_person.pk, other_branch.pk],
            )

        response = self.client.post(
            reverse("assignment_save", args=[self.week_start.isoformat(), self.person.pk, 1]),
            {
                "status": "work",
                "branch": self.branch.pk,
                "start_time": "08:00",
                "end_time": "16:00",
                "branch_filter": str(other_branch.pk),
            },
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Beto Ruiz")
        self.assertNotContains(response, "Ana Torres")

    def test_pages_include_accessibility_basics(self):
        schedule_response = self.client.get(
            reverse("schedule", args=[self.week_start.isoformat()])
        )
        self.assertContains(schedule_response, "Saltar al contenido")
        self.assertContains(schedule_response, 'id="page-content" tabindex="-1"')
        self.assertContains(schedule_response, "Grilla semanal de turnos")

        pending_response = self.client.get(
            reverse("schedule", args=[self.week_start.isoformat()]),
            {"filter": "pending"},
        )
        self.assertContains(pending_response, 'aria-current="true"')

        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()
        logs_response = self.client.get(reverse("logs"))
        self.assertContains(logs_response, "Historial de cambios")
        self.assertContains(logs_response, 'scope="col"')

    def test_completed_vacation_days_remain_marked_in_past_week(self):
        week_start = monday_of(date(2020, 1, 6))
        self.person.vacation_start = date(2020, 1, 7)
        self.person.vacation_end = date(2020, 1, 8)
        self.person.save()

        response = self.client.get(
            reverse("schedule", args=[week_start.isoformat()])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.content.count(b'class="shift-cell vacation-cell">Vacaciones</span>'),
            2,
        )

    def test_new_vacation_period_preserves_previous_vacation_cells(self):
        first_start = date(2020, 1, 7)
        first_end = date(2020, 1, 8)
        second_start = date(2020, 2, 4)
        second_end = date(2020, 2, 5)
        self.person.vacation_start = first_start
        self.person.vacation_end = first_end
        self.person.save()

        response = self.client.post(
            reverse("people"),
            {
                "person_id": self.person.pk,
                "name": self.person.name,
                "dni": "",
                "vacation_start": second_start.isoformat(),
                "vacation_end": second_end.isoformat(),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            set(
                VacationPeriod.objects.filter(person_id=self.person.pk).values_list(
                    "start_date", "end_date"
                )
            ),
            {(first_start, first_end), (second_start, second_end)},
        )
        for vacation_start in (first_start, second_start):
            week_start = monday_of(vacation_start)
            schedule_response = self.client.get(
                reverse("schedule", args=[week_start.isoformat()])
            )
            self.assertEqual(schedule_response.status_code, 200)
            self.assertEqual(
                schedule_response.content.count(
                    b'class="shift-cell vacation-cell">Vacaciones</span>'
                ),
                2,
            )

    def test_vacation_period_can_be_deleted_without_removing_other_periods(self):
        first_start = date(2020, 1, 7)
        first_end = date(2020, 1, 8)
        second_start = date(2020, 2, 4)
        second_end = date(2020, 2, 5)
        first_period = VacationPeriod.objects.create(
            person_id=self.person.pk,
            start_date=first_start,
            end_date=first_end,
        )
        second_period = VacationPeriod.objects.create(
            person_id=self.person.pk,
            start_date=second_start,
            end_date=second_end,
        )
        self.person.vacation_start = second_start
        self.person.vacation_end = second_end
        self.person.save()

        response = self.client.post(
            reverse(
                "vacation_period_delete",
                args=[self.person.pk, second_period.pk],
            ),
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(VacationPeriod.objects.filter(pk=first_period.pk).exists())
        self.assertFalse(VacationPeriod.objects.filter(pk=second_period.pk).exists())
        self.person.refresh_from_db()
        self.assertIsNone(self.person.vacation_start)
        self.assertIsNone(self.person.vacation_end)

        first_week = monday_of(first_start)
        response = self.client.get(
            reverse("schedule", args=[first_week.isoformat()])
        )
        self.assertEqual(
            response.content.count(
                b'class="shift-cell vacation-cell">Vacaciones</span>'
            ),
            2,
        )
        second_week = monday_of(second_start)
        response = self.client.get(
            reverse("schedule", args=[second_week.isoformat()])
        )
        self.assertEqual(
            response.content.count(
                b'class="shift-cell vacation-cell">Vacaciones</span>'
            ),
            0,
        )

    def test_rest_assignment_is_saved(self):
        response = self.client.post(
            reverse("assignment_save", args=[self.week_start.isoformat(), self.person.pk, 0]),
            {"status": "rest"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT branch_id, start_time, end_time FROM assignments "
                "WHERE week_start = %s AND person_id = %s AND day = 0",
                [self.week_start.isoformat(), self.person.pk],
            )
            self.assertEqual(cursor.fetchone(), (None, None, None))

    def test_incapacity_is_saved_displayed_and_counts_zero_worked_hours(self):
        response = self.client.post(
            reverse("assignment_save", args=[self.week_start.isoformat(), self.person.pk, 0]),
            {"status": "incapacity"},
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Incapacidad")
        assignment = week_assignments(self.week_start)[(self.person.pk, 0)]
        self.assertEqual(assignment["assignment_type"], "incapacity")
        self.assertIsNone(assignment["branch_id"])
        self.assertIsNone(assignment["start_time"])
        self.assertIsNone(assignment["end_time"])
        assignments = assignments_between(self.week_start, self.week_start)
        self.assertEqual(assignments[0][8], "incapacity")
        daily_metrics = assignment_period_metrics([self.person], assignments)
        summaries, _ = summarize_period([self.person], assignments, daily_metrics)
        self.assertEqual(daily_metrics, {})
        self.assertEqual(summaries[self.person.pk]["worked"], 0)
        workbook = load_workbook(
            BytesIO(fortnight_workbook(self.week_start, self.week_start, [self.person])),
            read_only=True,
        )
        self.assertEqual(workbook.active["C8"].value, "Incapacidad")
        self.assertEqual(workbook.active["F8"].value, 0)

        edit_response = self.client.get(
            reverse("schedule", args=[self.week_start.isoformat()]),
            {"edit": f"{self.person.pk}-0"},
        )
        self.assertEqual(edit_response.context["edit_form"]["status"].value(), "incapacity")

    def test_assignment_change_is_logged_with_ip_and_visible_in_logs(self):
        response = self.client.post(
            reverse("assignment_save", args=[self.week_start.isoformat(), self.person.pk, 0]),
            {
                "status": "work",
                "branch": self.branch.pk,
                "start_time": "08:00",
                "end_time": "16:00",
            },
            REMOTE_ADDR="203.0.113.7",
        )

        self.assertEqual(response.status_code, 200)
        entry = AuditLog.objects.get()
        self.assertEqual(entry.ip_address, "203.0.113.7")
        self.assertEqual(entry.action, "Turno guardado")
        self.assertIn("Ana Torres", entry.details)
        self.assertIn("Centro", entry.details)

        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()
        logs_response = self.client.get(reverse("logs"))
        self.assertEqual(logs_response.status_code, 200)
        self.assertContains(logs_response, "203.0.113.7")
        self.assertContains(logs_response, "Turno guardado")

    def test_assignment_update_logs_previous_and_new_shift(self):
        assignment_url = reverse(
            "assignment_save",
            args=[self.week_start.isoformat(), self.person.pk, 0],
        )
        initial_response = self.client.post(
            assignment_url,
            {
                "status": "work",
                "branch": self.branch.pk,
                "start_time": "08:00",
                "end_time": "16:00",
            },
        )
        self.assertEqual(initial_response.status_code, 200)

        updated_response = self.client.post(
            assignment_url,
            {
                "status": "work",
                "branch": self.branch.pk,
                "start_time": "09:00",
                "end_time": "17:00",
            },
            REMOTE_ADDR="203.0.113.8",
        )

        self.assertEqual(updated_response.status_code, 200)
        entry = AuditLog.objects.latest("pk")
        self.assertEqual(entry.ip_address, "203.0.113.8")
        self.assertEqual(entry.action, "Turno actualizado")
        self.assertIn("Antes: Centro · 08:00–16:00", entry.details)
        self.assertIn("Ahora: Centro · 09:00–17:00", entry.details)

    def test_unchanged_assignment_does_not_create_change_log(self):
        assignment_url = reverse(
            "assignment_save",
            args=[self.week_start.isoformat(), self.person.pk, 0],
        )
        payload = {
            "status": "work",
            "branch": self.branch.pk,
            "start_time": "08:00",
            "end_time": "16:00",
        }
        self.client.post(assignment_url, payload)
        self.client.post(assignment_url, payload)

        self.assertEqual(AuditLog.objects.count(), 1)

    def test_logs_require_analysis_session(self):
        response = self.client.get(reverse("logs"))

        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("analysis_unlock"), response["Location"])

    def test_invalid_assignment_is_not_logged(self):
        response = self.client.post(
            reverse("assignment_save", args=[self.week_start.isoformat(), self.person.pk, 0]),
            {"status": "work", "branch": self.branch.pk, "start_time": "", "end_time": ""},
            REMOTE_ADDR="203.0.113.7",
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(AuditLog.objects.exists())

    def test_people_page_uses_legacy_table(self):
        response = self.client.get(reverse("people"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ana Torres")

    def test_people_hide_vacations_by_default_and_can_be_shown(self):
        self.person.vacation_start = date(2026, 1, 5)
        self.person.vacation_end = date(2026, 1, 9)
        self.person.save()

        hidden = self.client.get(reverse("people"))
        self.assertContains(hidden, "Ocultar vacaciones")
        self.assertContains(hidden, 'value="1" checked')
        self.assertNotContains(hidden, "Vacaciones:")

        shown = self.client.get(reverse("people"), {"hide_vacations": "0"})
        self.assertContains(shown, "Vacaciones:")
        self.assertNotContains(shown, 'value="1" checked')

    def test_edit_person_form_loads_saved_vacation_dates(self):
        self.person.vacation_start = date(2026, 10, 12)
        self.person.vacation_end = date(2026, 10, 19)
        self.person.save()

        response = self.client.get(reverse("people"), {"edit": self.person.pk})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="person-form-card"', html=False)
        self.assertContains(response, 'name="vacation_start" value="2026-10-12"', html=False)
        self.assertContains(response, 'name="vacation_end" value="2026-10-19"', html=False)

    def test_branches_page_uses_legacy_table(self):
        response = self.client.get(reverse("branches"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Centro")

    def test_assignment_editor_renders(self):
        response = self.client.get(
            reverse("schedule", args=[self.week_start.isoformat()]),
            {"edit": f"{self.person.pk}-0"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Guardar turno")
        self.assertEqual(response.context["edit_form"]["status"].value(), "work")
        self.assertEqual(response.context["edit_form"]["branch"].value(), self.branch.pk)
        self.assertContains(response, '<dialog id="assignment-dialog"', html=False)
        self.assertContains(response, 'placeholder="12:00"', html=False)
        self.assertContains(response, 'placeholder="17:00"', html=False)
        self.assertContains(response, "data-split-shift-field hidden", html=False)

    def test_assignment_form_accepts_only_24_hour_hhmm(self):
        valid_form = AssignmentForm(data={
            "status": "work",
            "branch": self.branch.pk,
            "start_time": "12:00",
            "end_time": "17:00",
        })
        self.assertTrue(valid_form.is_valid())

        invalid_form = AssignmentForm(data={
            "status": "work",
            "branch": self.branch.pk,
            "start_time": "1:00",
            "end_time": "17:00",
        })
        self.assertFalse(invalid_form.is_valid())
        self.assertIn("Usa formato HH:MM.", invalid_form.errors["start_time"])

    def test_split_shift_form_requires_two_ordered_time_ranges(self):
        form = AssignmentForm(data={
            "status": "split_shift",
            "branch": self.branch.pk,
            "start_time": "10:00",
            "end_time": "14:00",
            "start_time_2": "17:00",
            "end_time_2": "18:30",
        })
        self.assertTrue(form.is_valid(), form.errors)

        missing_second_range = AssignmentForm(data={
            "status": "split_shift",
            "branch": self.branch.pk,
            "start_time": "10:00",
            "end_time": "14:00",
        })
        self.assertFalse(missing_second_range.is_valid())
        self.assertIn("start_time_2", missing_second_range.errors)
        self.assertIn("end_time_2", missing_second_range.errors)

        overlapping_ranges = AssignmentForm(data={
            "status": "split_shift",
            "branch": self.branch.pk,
            "start_time": "10:00",
            "end_time": "14:00",
            "start_time_2": "13:00",
            "end_time_2": "18:30",
        })
        self.assertFalse(overlapping_ranges.is_valid())
        self.assertIn("posterior a la primera salida", overlapping_ranges.errors["start_time_2"][0])

    def test_split_shift_saves_displays_and_excludes_break_from_analysis(self):
        week_start = monday_of(date(2026, 10, 7))
        assignment_url = reverse(
            "assignment_save",
            args=[week_start.isoformat(), self.person.pk, 0],
        )
        response = self.client.post(assignment_url, {
            "status": "split_shift",
            "branch": self.branch.pk,
            "start_time": "10:00",
            "end_time": "14:00",
            "start_time_2": "17:00",
            "end_time_2": "18:30",
        })
        self.assertEqual(response.status_code, 200)

        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT assignment_type, start_time, end_time, start_time_2, end_time_2 "
                "FROM assignments WHERE week_start = %s AND person_id = %s AND day = 0",
                [week_start.isoformat(), self.person.pk],
            )
            self.assertEqual(
                cursor.fetchone(),
                ("split_shift", "10:00", "14:00", "17:00", "18:30"),
            )

        schedule_response = self.client.get(
            reverse("schedule", args=[week_start.isoformat()])
        )
        self.assertContains(schedule_response, "10:00–14:00 / 17:00–18:30")
        split_shift_editor = self.client.get(
            reverse("schedule", args=[week_start.isoformat()]),
            {"edit": f"{self.person.pk}-0"},
        )
        self.assertContains(split_shift_editor, "data-split-shift-field>", html=False)
        self.assertNotContains(
            split_shift_editor, "data-split-shift-field hidden", html=False
        )

        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()
        analysis_response = self.client.get(
            reverse("analysis"),
            {"month": "2026-10", "fortnight": "1"},
        )
        self.assertEqual(analysis_response.status_code, 200)
        self.assertAlmostEqual(
            analysis_response.context["summaries"][self.person.pk]["worked"],
            5.5,
        )

        export_response = self.client.get(
            reverse("analysis_export"),
            {"month": "2026-10", "fortnight": "1"},
        )
        workbook = load_workbook(BytesIO(export_response.content), read_only=True)
        sheet = workbook.active
        self.assertAlmostEqual(sheet["C3"].value, 5.5)
        self.assertEqual(sheet["D8"].value, "10:00 / 17:00")
        self.assertEqual(sheet["E8"].value, "14:00 / 18:30")
        self.assertAlmostEqual(sheet["F8"].value, 5.5)

        schedule_export_response = self.client.get(
            reverse("schedule_export"),
            {"start": week_start.isoformat(), "end": week_start.isoformat()},
        )
        schedule_workbook = load_workbook(
            BytesIO(schedule_export_response.content), read_only=True
        )
        self.assertEqual(
            schedule_workbook.active["B3"].value,
            "Centro\n10:00–14:00 / 17:00–18:30",
        )

    @override_settings(ANALYSIS_PASSWORD="test-password")
    def test_analysis_page_renders_after_unlock(self):
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()
        response = self.client.get(
            reverse("analysis"),
            {"month": self.week_start.strftime("%Y-%m"), "fortnight": "1"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Horas trabajadas")
        self.assertContains(response, ">RN</th>")
        self.assertContains(response, "RN: recargo nocturno")
        self.assertContains(response, "Negativo: faltante contra la jornada de 7 h")
        self.assertContains(response, "Positivo: horas extra")
        self.assertContains(response, "Horas extra netas sobre 7 h trabajadas")
        self.assertContains(response, "turnos de 6 h o más")
        self.assertContains(response, "empiecen antes de 14:00")
        self.assertContains(response, "13:00–13:59")

    def test_novelties_requires_the_analysis_session(self):
        response = self.client.get(
            reverse("novelties_week", args=[self.week_start.isoformat()])
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            response["Location"],
            f"{reverse('analysis_unlock')}?next=%2Fnovedades%2F{self.week_start.isoformat()}%2F",
        )

    @override_settings(ANALYSIS_PASSWORD="test-password")
    def test_unlock_returns_to_requested_novelties_grid(self):
        target = reverse("novelties_week", args=[self.week_start.isoformat()])
        response = self.client.get(target)
        unlock_url = response["Location"]

        unlock_page = self.client.get(unlock_url)
        self.assertEqual(unlock_page.status_code, 200)
        self.assertContains(unlock_page, f'name="next" value="{target}"', html=False)

        response = self.client.post(
            unlock_url,
            {"password": "test-password", "next": target},
        )
        self.assertRedirects(response, target, fetch_redirect_response=False)
        grid = self.client.get(target)
        self.assertEqual(grid.status_code, 200)
        self.assertContains(grid, "schedule-grid-scroll")

    @override_settings(ANALYSIS_PASSWORD="test-password")
    def test_htmx_login_requests_full_page_reload_to_novelties(self):
        target = reverse("novelties_week", args=[self.week_start.isoformat()])
        unlock_url = f"{reverse('analysis_unlock')}?next=%2Fnovedades%2F{self.week_start.isoformat()}%2F"

        response = self.client.post(
            unlock_url,
            {"password": "test-password", "next": target},
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["HX-Redirect"], target)

    def test_novelties_menu_renders_grid_without_redirect(self):
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        response = self.client.get(reverse("novelties"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "schedule-grid-scroll")
        self.assertContains(response, "Novedades")
        self.assertContains(response, "MODO NOVEDADES")
        self.assertContains(response, "Ir a Turnos")
        self.assertContains(response, "Novedades activas")
        self.assertContains(response, 'title="Ana Torres"', html=False)
        self.assertNotContains(response, "Sin asignar")
        self.assertContains(response, "schedule-main-shell")

    def test_late_arrival_is_included_in_analysis_hours(self):
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        page = self.client.get(
            reverse("novelties_week", args=[self.week_start.isoformat()]),
            {"edit": f"{self.person.pk}-0"},
        )
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Novedades")
        self.assertContains(page, "Llegada tarde")
        self.assertContains(page, "No se presentó")
        self.assertContains(page, "Calamidad")
        self.assertContains(page, 'data-novelty-kind-select', html=False)
        self.assertContains(page, 'type="text"', html=False)
        self.assertContains(page, 'placeholder="00:00"', html=False)
        self.assertContains(page, 'pattern="(?:[01]\\d|2[0-3]):[0-5]\\d"', html=False)
        self.assertContains(page, "Observación (opcional)")

        response = self.client.post(
            reverse(
                "novelty_save",
                args=[self.week_start.isoformat(), self.person.pk, 0],
            ),
            {"kind": AssignmentNovelty.LATE_ARRIVAL, "actual_start_time": "09:30"},
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Llegada tarde")
        novelty = AssignmentNovelty.objects.get(
            week_start=self.week_start, person_id=self.person.pk, day=0
        )
        self.assertEqual(novelty.actual_start_time.strftime("%H:%M"), "09:30")
        assignments = assignments_between(self.week_start, self.week_start)
        self.assertEqual(assignments[0][5], "09:30")

        month = self.week_start.strftime("%Y-%m")
        fortnight = "1" if self.week_start.day <= 15 else "2"
        analysis = self.client.get(
            reverse("analysis"),
            {"month": month, "fortnight": fortnight},
        )
        self.assertEqual(analysis.status_code, 200)
        self.assertEqual(analysis.context["totals"]["worked"], 5.5)

        workbook_response = self.client.get(
            reverse("analysis_export"),
            {"month": month, "fortnight": fortnight},
        )
        workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
        self.assertEqual(workbook.active["C4"].value, 5.5)
        self.assertEqual(workbook.active["D8"].value, "09:30")

    def test_absence_and_calamity_do_not_add_worked_hours(self):
        with connection.cursor() as cursor:
            cursor.executemany(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, %s, %s, '08:00', '16:00', 'work')",
                [
                    (self.week_start.isoformat(), self.person.pk, day, self.branch.pk)
                    for day in (0, 1)
                ],
            )
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        for day, kind, observation, badge in (
            (0, AssignmentNovelty.DID_NOT_ATTEND, "No avisó", "No se presentó"),
            (1, AssignmentNovelty.CALAMITY, "Emergencia familiar", "Calamidad"),
        ):
            with self.subTest(kind=kind):
                response = self.client.post(
                    reverse(
                        "novelty_save",
                        args=[self.week_start.isoformat(), self.person.pk, day],
                    ),
                    {"kind": kind, "observation": observation},
                    HTTP_HX_REQUEST="true",
                )
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, badge)
                self.assertContains(response, observation)

                novelty = AssignmentNovelty.objects.get(
                    week_start=self.week_start, person_id=self.person.pk, day=day
                )
                self.assertIsNone(novelty.actual_start_time)
                self.assertEqual(novelty.observation, observation)

        assignments = assignments_between(self.week_start, self.week_start + timedelta(days=1))
        self.assertEqual(assignments[0][5:7], (None, AssignmentNovelty.DID_NOT_ATTEND))
        self.assertEqual(assignments[1][5:7], (None, AssignmentNovelty.CALAMITY))

        month = self.week_start.strftime("%Y-%m")
        fortnight = "1" if self.week_start.day <= 15 else "2"
        analysis = self.client.get(
            reverse("analysis"),
            {"month": month, "fortnight": fortnight},
        )
        metric = analysis.context["summaries"][self.person.pk]
        self.assertEqual(metric["days"], 0)
        self.assertEqual(metric["worked"], 0)

        workbook_response = self.client.get(
            reverse("analysis_export"),
            {"month": month, "fortnight": fortnight},
        )
        workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
        self.assertEqual(workbook.active["C3"].value, 0)
        self.assertEqual(workbook.active["F8"].value, 0)
        self.assertEqual(workbook.active["F9"].value, 0)
        self.assertIn("No se presentó", workbook.active["C8"].value)
        self.assertIn("Calamidad", workbook.active["C9"].value)
        self.assertEqual(workbook.active["K8"].value, "No avisó")
        self.assertEqual(workbook.active["K9"].value, "Emergencia familiar")

    def test_changing_shift_times_preserves_existing_novelty(self):
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            actual_start_time="09:30",
        )

        response = self.client.post(
            reverse(
                "assignment_save",
                args=[self.week_start.isoformat(), self.person.pk, 0],
            ),
            {
                "status": "work",
                "branch": self.branch.pk,
                "start_time": "09:00",
                "end_time": "17:00",
            },
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        novelty = AssignmentNovelty.objects.get()
        self.assertIsNone(novelty.archived_at)
        self.assertEqual(novelty.actual_start_time.strftime("%H:%M"), "09:30")
        assignments = assignments_between(self.week_start, self.week_start)
        self.assertEqual(assignments[0][5], "09:30")
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        novelties_page = self.client.get(
            reverse("novelties_week", args=[self.week_start.isoformat()])
        )

        self.assertEqual(novelties_page.status_code, 200)
        self.assertContains(novelties_page, "Llegada tarde")
        self.assertNotContains(novelties_page, "Histórico ·")

    def test_changing_shift_branch_archives_outdated_novelty(self):
        other_branch = Branch.objects.create(name="Norte", color="#2F6FED")
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            actual_start_time="09:30",
        )

        response = self.client.post(
            reverse(
                "assignment_save",
                args=[self.week_start.isoformat(), self.person.pk, 0],
            ),
            {
                "status": "work",
                "branch": other_branch.pk,
                "start_time": "08:00",
                "end_time": "16:00",
            },
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        novelty = AssignmentNovelty.objects.get()
        self.assertIsNotNone(novelty.archived_at)
        self.assertEqual(novelty.actual_start_time.strftime("%H:%M"), "09:30")

    def test_removing_shift_archives_associated_novelty(self):
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            actual_start_time="09:30",
        )
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        response = self.client.post(
            reverse(
                "assignment_save",
                args=[self.week_start.isoformat(), self.person.pk, 0],
            ),
            {"status": "unassigned"},
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        novelty = AssignmentNovelty.objects.get()
        self.assertIsNotNone(novelty.archived_at)
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COUNT(*) FROM assignments "
                "WHERE week_start = %s AND person_id = %s AND day = 0",
                [self.week_start.isoformat(), self.person.pk],
            )
            self.assertEqual(cursor.fetchone()[0], 0)

        history_response = self.client.get(
            reverse("novelties_week", args=[self.week_start.isoformat()])
        )
        self.assertEqual(history_response.status_code, 200)
        self.assertContains(history_response, "Histórico · Llegada tarde")
        self.assertContains(history_response, "09:30")

    def test_selecting_no_novelty_archives_record_without_changing_shift(self):
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            kind=AssignmentNovelty.CALAMITY,
            observation="Registro equivocado",
        )
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        response = self.client.post(
            reverse(
                "novelty_save",
                args=[self.week_start.isoformat(), self.person.pk, 0],
            ),
            {"kind": "", "actual_start_time": "", "observation": ""},
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Novedad conservada en el historial.")
        self.assertNotContains(response, "+ Calamidad")
        novelty = AssignmentNovelty.objects.get()
        self.assertIsNotNone(novelty.archived_at)
        self.assertEqual(novelty.observation, "Registro equivocado")
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT branch_id, start_time, end_time FROM assignments "
                "WHERE week_start = %s AND person_id = %s AND day = 0",
                [self.week_start.isoformat(), self.person.pk],
            )
            self.assertEqual(cursor.fetchone(), (self.branch.pk, "08:00", "16:00"))

    def test_clearing_week_archives_novelties(self):
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            actual_start_time="09:30",
        )

        clear_week(self.week_start)

        novelty = AssignmentNovelty.objects.get()
        self.assertIsNotNone(novelty.archived_at)

    def test_deleting_branch_archives_associated_novelties(self):
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            actual_start_time="09:30",
        )

        delete_branch(self.branch)

        novelty = AssignmentNovelty.objects.get()
        self.assertIsNotNone(novelty.archived_at)

    def test_history_can_be_deleted_for_one_cell_without_removing_current_novelty(self):
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, 0, %s, '08:00', '16:00', 'work')",
                [self.week_start.isoformat(), self.person.pk, self.branch.pk],
            )
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            actual_start_time="09:30",
            archived_at=timezone.now(),
        )
        current_novelty = AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            actual_start_time="10:00",
        )
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        page = self.client.get(
            reverse("novelties_week", args=[self.week_start.isoformat()])
        )
        self.assertContains(page, "Borrar historial")

        response = self.client.post(
            reverse(
                "novelty_history_delete",
                args=[self.week_start.isoformat(), self.person.pk, 0],
            ),
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Historial de la celda eliminado.")
        self.assertFalse(
            AssignmentNovelty.objects.filter(archived_at__isnull=False).exists()
        )
        self.assertEqual(
            AssignmentNovelty.objects.get(archived_at__isnull=True).pk,
            current_novelty.pk,
        )
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT branch_id, start_time, end_time FROM assignments "
                "WHERE week_start = %s AND person_id = %s AND day = 0",
                [self.week_start.isoformat(), self.person.pk],
            )
            self.assertEqual(cursor.fetchone(), (self.branch.pk, "08:00", "16:00"))

    def test_novelty_time_must_fall_after_shift_start(self):
        no_attendance_form = LateArrivalForm(
            data={
                "kind": AssignmentNovelty.DID_NOT_ATTEND,
                "actual_start_time": "",
                "observation": "No avisó",
            },
            start_time="08:00",
            end_time="16:00",
        )
        self.assertTrue(no_attendance_form.is_valid(), no_attendance_form.errors)

        invalid_format_form = LateArrivalForm(
            data={"kind": AssignmentNovelty.LATE_ARRIVAL, "actual_start_time": "9:30"},
            start_time="08:00",
            end_time="16:00",
        )
        self.assertFalse(invalid_format_form.is_valid())
        self.assertEqual(
            invalid_format_form.errors["actual_start_time"],
            ["Usa formato HH:MM, por ejemplo 09:30."],
        )

        form = LateArrivalForm(
            data={"kind": AssignmentNovelty.LATE_ARRIVAL, "actual_start_time": "07:30"},
            start_time="08:00",
            end_time="16:00",
        )
        self.assertFalse(form.is_valid())
        self.assertIn("posterior al inicio del turno", form.errors["actual_start_time"][0])

        overnight_form = LateArrivalForm(
            data={"kind": AssignmentNovelty.LATE_ARRIVAL, "actual_start_time": "00:30"},
            start_time="22:00",
            end_time="06:00",
        )
        self.assertTrue(overnight_form.is_valid())

        split_shift_times = {
            "kind": AssignmentNovelty.LATE_ARRIVAL,
            "actual_start_time": "",
            "actual_start_time_2": "",
        }
        split_shift_form = LateArrivalForm(
            data=split_shift_times,
            start_time="10:00",
            end_time="14:00",
            start_time_2="17:00",
            end_time_2="18:30",
            split_shift=True,
        )
        self.assertFalse(split_shift_form.is_valid())
        self.assertIn(
            "Indica al menos una de las dos horas reales de llegada.",
            split_shift_form.non_field_errors(),
        )

        first_arrival_only = LateArrivalForm(
            data={**split_shift_times, "actual_start_time": "10:30"},
            start_time="10:00",
            end_time="14:00",
            start_time_2="17:00",
            end_time_2="18:30",
            split_shift=True,
        )
        self.assertTrue(first_arrival_only.is_valid(), first_arrival_only.errors)

        second_arrival_only = LateArrivalForm(
            data={**split_shift_times, "actual_start_time_2": "17:30"},
            start_time="10:00",
            end_time="14:00",
            start_time_2="17:00",
            end_time_2="18:30",
            split_shift=True,
        )
        self.assertTrue(second_arrival_only.is_valid(), second_arrival_only.errors)

    def test_split_shift_novelty_records_and_applies_both_arrival_times(self):
        week_start = monday_of(date(2026, 10, 7))
        assignment_url = reverse(
            "assignment_save",
            args=[week_start.isoformat(), self.person.pk, 0],
        )
        self.client.post(assignment_url, {
            "status": "split_shift",
            "branch": self.branch.pk,
            "start_time": "10:00",
            "end_time": "14:00",
            "start_time_2": "17:00",
            "end_time_2": "18:30",
        })
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        novelty_page = self.client.get(
            reverse("novelties_week", args=[week_start.isoformat()]),
            {"edit": f"{self.person.pk}-0"},
        )
        self.assertEqual(novelty_page.status_code, 200)
        self.assertContains(novelty_page, 'name="actual_start_time"')
        self.assertContains(novelty_page, 'name="actual_start_time_2"')
        self.assertContains(novelty_page, "Llegada real 1")
        self.assertContains(novelty_page, "Llegada real 2")

        response = self.client.post(
            reverse(
                "novelty_save",
                args=[week_start.isoformat(), self.person.pk, 0],
            ),
            {
                "kind": AssignmentNovelty.LATE_ARRIVAL,
                "actual_start_time": "10:30",
                "actual_start_time_2": "17:30",
                "observation": "",
            },
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        novelty = AssignmentNovelty.objects.get()
        self.assertEqual(novelty.actual_start_time.strftime("%H:%M"), "10:30")
        self.assertEqual(novelty.actual_start_time_2.strftime("%H:%M"), "17:30")

        assignments = assignments_between(week_start, week_start)
        self.assertEqual(assignments[0][5], "10:30")
        self.assertEqual(assignments[0][11], "17:30")
        summary, _ = summarize_period([self.person], assignments)
        self.assertAlmostEqual(summary[self.person.pk]["worked"], 4.5)

        normal_assignment_url = reverse(
            "assignment_save",
            args=[week_start.isoformat(), self.person.pk, 1],
        )
        self.client.post(normal_assignment_url, {
            "status": "work",
            "branch": self.branch.pk,
            "start_time": "08:00",
            "end_time": "16:00",
        })
        normal_novelty_page = self.client.get(
            reverse("novelties_week", args=[week_start.isoformat()]),
            {"edit": f"{self.person.pk}-1"},
        )
        self.assertNotContains(normal_novelty_page, 'name="actual_start_time_2"')

    def test_schedule_export_downloads_excel(self):
        response = self.client.get(
            reverse("schedule_export"),
            {"start": self.week_start.isoformat(), "end": self.week_start.isoformat()},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_incapacity_appears_in_both_excel_exports(self):
        assignment_response = self.client.post(
            reverse("assignment_save", args=[self.week_start.isoformat(), self.person.pk, 0]),
            {"status": "incapacity"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(assignment_response.status_code, 200)

        schedule_response = self.client.get(
            reverse("schedule_export"),
            {"start": self.week_start.isoformat(), "end": self.week_start.isoformat()},
        )
        self.assertEqual(schedule_response.status_code, 200)
        schedule_sheet = load_workbook(
            BytesIO(schedule_response.content), read_only=True
        ).active
        self.assertEqual(schedule_sheet["B3"].value, "Incapacidad")
        self.assertEqual(schedule_sheet["C3"].value, 0)

        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()
        month = self.week_start.strftime("%Y-%m")
        fortnight = "1" if self.week_start.day <= 15 else "2"
        analysis_response = self.client.get(
            reverse("analysis_export"),
            {"month": month, "fortnight": fortnight},
        )
        self.assertEqual(analysis_response.status_code, 200)
        analysis_sheet = load_workbook(
            BytesIO(analysis_response.content), read_only=True
        ).active
        self.assertEqual(analysis_sheet["C8"].value, "Incapacidad")
        self.assertEqual(analysis_sheet["F8"].value, 0)

    def test_schedule_export_marks_all_vacation_periods(self):
        VacationPeriod.objects.create(
            person_id=self.person.pk,
            start_date=date(2020, 1, 7),
            end_date=date(2020, 1, 8),
        )
        VacationPeriod.objects.create(
            person_id=self.person.pk,
            start_date=date(2020, 2, 4),
            end_date=date(2020, 2, 5),
        )

        response = self.client.get(
            reverse("schedule_export"),
            {"start": "2020-01-06", "end": "2020-02-06"},
        )

        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), read_only=True).active
        self.assertEqual(sheet["C3"].value, "Vacaciones")
        self.assertEqual(sheet["D3"].value, "Vacaciones")
        self.assertEqual(sheet["AE3"].value, "Vacaciones")
        self.assertEqual(sheet["AF3"].value, "Vacaciones")

    def test_overnight_shift_metrics(self):
        self.assertEqual(
            shift_metrics("22:00", "06:00"), (8.0, 1.0, 0.0, 1.0, 8.0)
        )

    def test_lunch_is_deducted_only_when_shift_touches_it(self):
        for start_time, end_time, expected in (
            ("09:00", "17:00", (7.0, 0.0, 0.0, 0.0, 0.0)),
            ("09:00", "18:00", (8.0, 1.0, 1.0, 0.0, 0.0)),
            ("10:30", "19:00", (7.5, 0.5, 0.5, 0.0, 0.0)),
            ("13:30", "22:00", (7.5, 0.5, 0.0, 0.5, 3.0)),
            ("12:00", "20:00", (7.0, 0.0, 0.0, 0.0, 1.0)),
            ("13:00", "22:00", (8.0, 1.0, 0.0, 1.0, 3.0)),
        ):
            with self.subTest(start_time=start_time, end_time=end_time):
                self.assertEqual(shift_metrics(start_time, end_time), expected)

    def test_lunch_is_not_deducted_outside_the_lunch_window(self):
        for start_time, end_time, expected in (
            ("05:00", "13:00", (8.0, 1.0, 1.0, 0.0, 1.0)),
            ("14:00", "22:00", (8.0, 1.0, 0.0, 1.0, 3.0)),
            ("22:00", "06:00", (8.0, 1.0, 0.0, 1.0, 8.0)),
        ):
            with self.subTest(start_time=start_time, end_time=end_time):
                self.assertEqual(shift_metrics(start_time, end_time), expected)

    def test_lunch_requires_a_six_hour_shift(self):
        self.assertEqual(
            shift_metrics("10:00", "16:00"), (5.0, 0.0, 0.0, 0.0, 0.0)
        )
        self.assertEqual(
            shift_metrics("10:00", "15:30"), (5.5, 0.0, 0.0, 0.0, 0.0)
        )

    def test_lunch_window_is_13_00_to_13_59(self):
        self.assertFalse(shift_touches_lunch("05:00", "13:00"))
        self.assertTrue(shift_touches_lunch("05:00", "13:01"))
        self.assertTrue(shift_touches_lunch("13:00", "20:00"))
        self.assertTrue(shift_touches_lunch("13:59", "22:00"))
        self.assertFalse(shift_touches_lunch("14:00", "22:00"))

    def test_shifts_starting_at_or_after_14_00_skip_lunch(self):
        for start_time, end_time in (
            ("14:00", "22:00"),
            ("14:01", "22:00"),
            ("15:00", "23:00"),
            ("22:00", "14:00"),
            ("23:00", "13:30"),
        ):
            with self.subTest(start_time=start_time, end_time=end_time):
                self.assertFalse(shift_touches_lunch(start_time, end_time))
        self.assertTrue(shift_touches_lunch("05:00", "14:00"))
        self.assertTrue(shift_touches_lunch("13:59", "22:00"))

    def test_lunch_does_not_apply_to_split_shifts(self):
        self.assertEqual(
            shift_segments_metrics(
                (("09:00", "13:00"), ("14:00", "18:30"))
            ),
            (8.5, 1.5, 1.5, 0.0, 0.0),
        )
        self.assertEqual(
            shift_segments_metrics(
                (("09:00", "14:00"), ("15:00", "19:00"))
            ),
            (9.0, 2.0, 2.0, 0.0, 0.0),
        )

    def test_overtime_starts_after_the_lunch_hour(self):
        self.assertEqual(
            shift_metrics("09:00", "19:00"), (9.0, 2.0, 2.0, 0.0, 0.0)
        )

    def test_night_surcharge_counts_all_night_hours(self):
        for start_time, end_time, expected in (
            ("08:00", "19:30", 0.5),
            ("14:00", "20:00", 1.0),
            ("09:30", "20:00", 1.0),
            ("05:00", "13:00", 1.0),
            ("09:00", "18:00", 0.0),
            ("19:00", "22:00", 3.0),
        ):
            with self.subTest(start_time=start_time, end_time=end_time):
                self.assertEqual(
                    shift_metrics(start_time, end_time)[4], expected
                )

    def test_night_surcharge_adds_both_segments_of_a_split_shift(self):
        self.assertEqual(
            shift_segments_metrics(
                (("18:00", "22:00"), ("23:00", "02:00"))
            )[4],
            6.0,
        )

    def test_late_arrival_after_lunch_skips_the_deduction(self):
        assignments = [(
            self.week_start,
            self.person.pk,
            self.branch.pk,
            "09:00",
            "18:00",
            "14:00",
            None,
            "",
            "work",
            None,
            None,
            None,
        )]
        daily_metrics = assignment_period_metrics([self.person], assignments)
        self.assertEqual(daily_metrics[0]["worked"], 4 * 60)
        self.assertEqual(daily_metrics[0]["night_surcharge"], 0)

    def test_analysis_preserves_negative_daytime_balance_and_nighttime_overtime(self):
        week_start = date(2026, 10, 5)
        shift_data = (
            ("09:30", "18:30", None, None),
            ("09:30", "14:00", None, None),
            ("09:30", "18:00", None, None),
            ("09:30", "20:00", "09:37", AssignmentNovelty.LATE_ARRIVAL),
            ("09:30", "15:00", None, None),
            ("09:30", "18:30", None, None),
        )
        assignments = [
            (
                week_start + timedelta(days=day),
                self.person.pk,
                self.branch.pk,
                start_time,
                end_time,
                actual_start_time,
                novelty_kind,
                "",
            )
            for day, (start_time, end_time, actual_start_time, novelty_kind) in zip(
                (0, 1, 2, 3, 4, 6), shift_data
            )
        ]

        daily_metrics = assignment_period_metrics([self.person], assignments)
        summaries, _ = summarize_period([self.person], assignments, daily_metrics)
        metric = summaries[self.person.pk]

        self.assertAlmostEqual(metric["worked"], 42 + 53 / 60)
        self.assertAlmostEqual(metric["overtime"], 53 / 60)
        self.assertAlmostEqual(metric["daytime"], -7 / 60)
        self.assertEqual(metric["nighttime"], 1)
        self.assertEqual(daily_metrics[3]["nighttime"], 60)
        self.assertEqual(daily_metrics[1]["daytime"], -150)
        self.assertEqual(daily_metrics[4]["daytime"], -90)
        self.assertEqual(
            sum(item["overtime"] for item in daily_metrics.values()),
            53,
        )

    def test_analysis_nets_overtime_against_shorter_workdays(self):
        week_start = monday_of(date(2026, 10, 8))
        shift_end_times = ("17:16", "17:16", "17:16", "17:16", "13:00", "15:15", "15:00")
        with connection.cursor() as cursor:
            cursor.executemany(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, %s, %s, '08:00', %s, 'work')",
                [
                    (week_start.isoformat(), self.person.pk, day, self.branch.pk, end_time)
                    for day, end_time in enumerate(shift_end_times)
                ],
            )

        assignments = assignments_between(
            week_start, week_start + timedelta(days=6)
        )
        summaries, totals = summarize_period([self.person], assignments)
        metric = summaries[self.person.pk]

        self.assertEqual(metric["days"], 7)
        self.assertAlmostEqual(metric["worked"], 3019 / 60)
        self.assertAlmostEqual(metric["overtime"], 79 / 60)
        self.assertAlmostEqual(metric["daytime"], 79 / 60)
        self.assertEqual(metric["nighttime"], 0)
        self.assertAlmostEqual(totals["overtime"], 79 / 60)

        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()
        month = week_start.strftime("%Y-%m")
        fortnight = "1" if week_start.day <= 15 else "2"
        response = self.client.get(
            reverse("analysis"),
            {"month": month, "fortnight": fortnight},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="metric-positive"')
        self.assertAlmostEqual(
            response.context["summaries"][self.person.pk]["overtime"],
            79 / 60,
        )

        workbook_response = self.client.get(
            reverse("analysis_export"),
            {"month": month, "fortnight": fortnight},
        )
        workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
        self.assertEqual(
            tuple(workbook.active.cell(row=2, column=column).value for column in range(1, 8)),
            (
                "Nombre", "Días trabajados", "Horas trabajadas",
                "Horas diurnas", "Horas nocturnas", "Horas extra", "RN",
            ),
        )
        self.assertAlmostEqual(workbook.active["C3"].value, 3019 / 60)
        self.assertAlmostEqual(workbook.active["D3"].value, 79 / 60)
        self.assertEqual(workbook.active["E3"].value, 0)
        self.assertAlmostEqual(workbook.active["F3"].value, 79 / 60)
        self.assertEqual(workbook.active["G3"].value, 0)
        self.assertEqual(workbook.active["G7"].value, "Horas extra")
        self.assertEqual(workbook.active["H7"].value, "HD")
        self.assertEqual(workbook.active["I7"].value, "HN")
        self.assertEqual(workbook.active["J7"].value, "RN")
        self.assertAlmostEqual(
            sum(workbook.active.cell(row=row, column=7).value for row in range(8, 15)),
            workbook.active["F3"].value,
        )
        self.assertAlmostEqual(
            sum(workbook.active.cell(row=row, column=8).value for row in range(8, 15)),
            workbook.active["D3"].value,
        )
        self.assertAlmostEqual(
            sum(workbook.active.cell(row=row, column=9).value for row in range(8, 15)),
            workbook.active["E3"].value,
        )
        self.assertAlmostEqual(
            sum(workbook.active.cell(row=row, column=10).value for row in range(8, 15)),
            workbook.active["G3"].value,
        )

    def test_analysis_export_shows_daily_daytime_and_nighttime_overtime(self):
        with connection.cursor() as cursor:
            cursor.executemany(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, %s, %s, %s, %s, 'work')",
                [
                    (
                        self.week_start.isoformat(), self.person.pk, 0,
                        self.branch.pk, "22:00", "08:00",
                    ),
                    (
                        self.week_start.isoformat(), self.person.pk, 1,
                        self.branch.pk, "15:00", "01:00",
                    ),
                ],
            )
        session = self.client.session
        session["analysis_authenticated"] = True
        session.save()

        workbook_response = self.client.get(
            reverse("analysis_export"),
            {
                "month": self.week_start.strftime("%Y-%m"),
                "fortnight": "1" if self.week_start.day <= 15 else "2",
            },
        )

        self.assertEqual(workbook_response.status_code, 200)
        workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
        sheet = workbook.active
        self.assertEqual(
            tuple(sheet.cell(row=2, column=column).value for column in range(1, 8)),
            (
                "Nombre", "Días trabajados", "Horas trabajadas",
                "Horas diurnas", "Horas nocturnas", "Horas extra", "RN",
            ),
        )
        self.assertEqual((sheet["D3"].value, sheet["E3"].value, sheet["F3"].value), (2, 4, 6))
        self.assertEqual(sheet["G3"].value, 14)
        self.assertEqual((sheet["G8"].value, sheet["H8"].value, sheet["I8"].value), (3, 2, 1))
        self.assertEqual((sheet["G9"].value, sheet["H9"].value, sheet["I9"].value), (3, 0, 3))
        self.assertEqual((sheet["J8"].value, sheet["J9"].value), (8, 6))
        self.assertEqual(sheet["F3"].value, sheet["G8"].value + sheet["G9"].value)
        self.assertEqual(sheet["D3"].value, sheet["H8"].value + sheet["H9"].value)
        self.assertEqual(sheet["E3"].value, sheet["I8"].value + sheet["I9"].value)
        self.assertEqual(sheet["G3"].value, sheet["J8"].value + sheet["J9"].value)

    def test_absence_and_calamity_reduce_overtime_and_daily_quota(self):
        shift_end_times = ("17:16", "17:16", "17:16", "17:16", "13:00", "15:15", "15:00")
        with connection.cursor() as cursor:
            cursor.executemany(
                "INSERT INTO assignments (week_start, person_id, day, branch_id, start_time, end_time, assignment_type) "
                "VALUES (%s, %s, %s, %s, '08:00', %s, 'work')",
                [
                    (self.week_start.isoformat(), self.person.pk, day, self.branch.pk, end_time)
                    for day, end_time in enumerate(shift_end_times)
                ],
            )
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=0,
            kind=AssignmentNovelty.DID_NOT_ATTEND,
        )
        AssignmentNovelty.objects.create(
            week_start=self.week_start,
            person_id=self.person.pk,
            day=4,
            kind=AssignmentNovelty.CALAMITY,
        )

        assignments = assignments_between(
            self.week_start, self.week_start + timedelta(days=6)
        )
        summaries, _ = summarize_period([self.person], assignments)
        metric = summaries[self.person.pk]

        self.assertEqual(metric["days"], 5)
        self.assertAlmostEqual(metric["worked"], 2223 / 60)
        self.assertAlmostEqual(metric["overtime"], 123 / 60)
        self.assertAlmostEqual(metric["daytime"], 123 / 60)
        self.assertEqual(metric["nighttime"], 0)


class WeekDateTests(TestCase):
    def test_monday_of_week(self):
        self.assertEqual(monday_of(date(2026, 10, 3)), date(2026, 9, 28))

    def test_capitalize_person_name_and_keep_dni_acronym(self):
        self.assertEqual(capitalize("ANA TORRES"), "Ana Torres")
        self.assertEqual(capitalize("ANA TORRES · DNI 123"), "Ana Torres · DNI 123")


class TemplateTagTests(TestCase):
    def test_signed_class_marks_negative_and_positive_values(self):
        self.assertEqual(signed_class(-1.5), "metric-negative")
        self.assertEqual(signed_class(0), "")
        self.assertEqual(signed_class(2), "metric-positive")
        self.assertEqual(signed_class(None), "")
